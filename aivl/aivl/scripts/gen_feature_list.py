# -*- coding: utf-8 -*-
"""生成 AI 投标报价系统功能清单明细 Excel（含工时预估）。"""
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "release" / "AI投标报价系统-功能清单明细.xlsx"
OUT_ALT = ROOT / "release" / "AI投标报价系统-功能清单明细-一人月分摊.xlsx"

# 工时口径：一人一个月 = 双休、每天 8 小时、约 22 工作日 = 176 小时，按复杂度权重分摊到各功能
# WEIGHTS 为相对权重；实际写入 Excel 的小时数会归一化到 MONTH_HOURS
MONTH_WORKDAYS = 22
HOURS_PER_DAY = 8
MONTH_HOURS = MONTH_WORKDAYS * HOURS_PER_DAY  # 176

# 序号, 模块, 功能名称, 功能说明, 入口/页面或API, 优先级, 状态, 复杂度权重, 备注
ROWS = [
    (1, "系统与登录", "用户登录", "账号密码登录，签发 Bearer Token 并持久化到本地", "登录页；POST /api/auth/login", "核心", "已实现", 12, ""),
    (2, "系统与登录", "登录态校验", "进入系统时校验 Token；过期触发重新登录", "App 启动；GET /api/auth/me", "核心", "已实现", 6, ""),
    (3, "系统与登录", "退出登录", "注销会话并清除本地 Token", "顶栏退出；POST /api/auth/logout", "重要", "已实现", 2, ""),
    (4, "系统与登录", "API 鉴权中间件", "除健康检查/登录外，业务 API 需登录", "main.py 中间件", "核心", "已实现", 8, ""),
    (5, "系统与登录", "服务健康检查", "返回服务状态、LLM/邮件是否已配置", "GET /api/health", "一般", "已实现", 2, ""),
    (6, "项目管理", "项目列表", "左侧展示全部项目及阶段状态标签", "侧栏；GET /api/projects", "核心", "已实现", 8, ""),
    (7, "项目管理", "新建项目", "填写名称/备注创建项目", "新建项目；POST /api/projects", "核心", "已实现", 6, ""),
    (8, "项目管理", "查看项目详情", "加载项目信息、文档树、公告分析结果等", "选中项目；GET /api/projects/{id}", "核心", "已实现", 12, ""),
    (9, "项目管理", "删除项目", "删除当前项目及相关数据", "删除项目；DELETE /api/projects/{id}", "重要", "已实现", 4, ""),
    (10, "项目管理", "分阶段工作流导航", "四阶段：上传资料→招标公告→投标须知→投标报价", "顶部 Steps", "核心", "已实现", 16, ""),
    (11, "文件入库", "上传招标 ZIP", "上传完整招标包，后台解压并索引文档", "阶段一；POST .../upload-zip", "核心", "已实现", 24, "大包后台解压"),
    (12, "文件入库", "ZIP 目录树预览", "树形展示包内结构，标注公告区/附件", "阶段一目录树", "重要", "已实现", 12, ""),
    (13, "文件入库", "自动定位招标公告", "入库后识别招标通告等公告文档路径", "详情 notice_doc", "核心", "已实现", 8, ""),
    (14, "文件入库", "文档角色分类", "按路径/文件名区分公告、须知、附件等", "ingest 入库索引", "重要", "已实现", 12, ""),
    (15, "招标公告", "分析招标公告", "LLM 抽取结构化字段、原子要求、待确认问题", "阶段二；POST .../analyze-notice", "核心", "已实现", 32, "依赖 DASHSCOPE_API_KEY"),
    (16, "招标公告", "分析模板模式", "支持基础模板/通用 LLM；项目可锁定模式", "分析弹窗 template_mode", "重要", "已实现", 8, ""),
    (17, "招标公告", "关注关键词检索", "可选关键词，结果含 focus_hits", "分析弹窗", "重要", "已实现", 6, ""),
    (18, "招标公告", "查看公告分析结果", "分章节字段表、原子要求、开放问题、来源摘要", "阶段二；GET .../notice-analysis", "核心", "已实现", 16, ""),
    (19, "招标公告", "公告详情数据", "按解读表模板全量抽取多 Sheet 原子需求", "公告详情数据；GET .../structured-export", "核心", "已实现", 40, "分块并行抽取，权重最高"),
    (20, "招标公告", "导出结构化解读 Excel", "按《招标文件结构化解读表》生成 xlsx", "GET .../export-excel", "重要", "已实现", 16, "API 已实现，主 UI 下载入口较弱"),
    (21, "投标须知", "动态扫描附件清单", "从 ZIP 须知目录识别附件01–09及正文", "重新扫描；POST .../instructions/scan", "核心", "已实现", 20, ""),
    (22, "投标须知", "读取须知状态", "获取工作流步骤、附件列表、条件与基础数据", "GET .../instructions", "核心", "已实现", 8, ""),
    (23, "投标须知", "分析投标须知正文", "LLM 结构化摘要：字段、硬性规则、附件提示、风险", "分析投标须知；POST .../analyze-body", "核心", "已实现", 24, ""),
    (24, "投标须知", "须知关注关键词", "分析时指定关键词并展示命中", "须知分析弹窗", "一般", "已实现", 4, ""),
    (25, "投标须知", "须知子流程导航", "识别清单→基础数据→分级编制→联动校验→完成", "阶段三内 Steps", "核心", "已实现", 12, ""),
    (26, "投标须知", "更新须知工作流状态", "保存步骤、条件、基础数据、附件状态；可校验/完成", "PATCH .../instructions", "核心", "已实现", 16, ""),
    (27, "投标须知", "统一基础数据维护", "项目/招标方/投标人/法人/代理人/截标/币种等", "须知子步骤②", "核心", "已实现", 20, ""),
    (28, "投标须知", "导入文件回填基础数据", "Word/PDF/Excel/TXT 分析后写入基础字段", "POST .../base-data/import", "重要", "已实现", 16, ""),
    (29, "投标须知", "条件判断", "记录条件决策，驱动附件是否纳入/无需提交", "条件表单", "重要", "已实现", 12, ""),
    (30, "投标须知", "一致性校验", "名称/金额/人员/进度等联动校验", "子步骤④ run_validate", "核心", "已实现", 20, ""),
    (31, "投标须知", "完成须知阶段", "校验通过后完成本阶段，可进入报价", "子步骤⑤完成", "核心", "已实现", 4, ""),
    (32, "附件编制", "企业/项目主数据", "维护企业、法人、代理人、项目信息供附件预填", "GET/PUT .../org-master", "重要", "已实现", 12, ""),
    (33, "附件编制", "打开附件工作台", "按附件编码加载字段、模板、状态、依赖", "GET .../attachments/{code}/workbench", "核心", "已实现", 24, ""),
    (34, "附件编制", "AI 分析附件填空字段", "从模板正文识别待填字段", "分析字段；POST .../analyze-fields", "重要", "已实现", 16, ""),
    (35, "附件编制", "保存附件字段", "手工编辑后保存字段值", "PUT .../fields", "核心", "已实现", 6, ""),
    (36, "附件编制", "导入资料回填附件字段", "导入文件分析回填（组织图/人员表除外）", "POST .../import-fill", "重要", "已实现", 20, ""),
    (37, "附件编制", "完成本附件", "标记附件完成并推进进度", "POST .../complete", "核心", "已实现", 4, ""),
    (38, "附件编制", "导出填写后 Word", "基于原模板套打导出待盖章 docx", "导出 Word；GET .../export-docx", "核心", "已实现", 32, "含版式/下划线处理联调"),
    (39, "附件编制", "附件01 授权委托书", "已备案/未备案分支；主数据预填；模板导出", "A01 工作台", "核心", "已实现", 24, ""),
    (40, "附件编制", "附件02 投标承诺函", "锁定报价；大小写金额校验；现场考察确认", "A02 工作台", "核心", "已实现", 24, ""),
    (41, "附件编制", "附件04 人员配备与履历表", "增删改人员行；导入合并/追加/替换", "A04；personnel 接口", "核心", "已实现", 32, ""),
    (42, "附件编制", "附件03 组织机构图", "由 A04 人员生成架构图，可编辑并确认", "A03；org-chart 接口", "核心", "已实现", 28, ""),
    (43, "附件编制", "附件05 劳动力计划与曲线", "分阶段人数、曲线 SVG；可生成草稿", "plan-workbench labor_plan", "重要", "已实现", 20, ""),
    (44, "附件编制", "附件06 机械设备配备计划", "设备机具表格维护与草稿生成", "plan-workbench equipment_plan", "重要", "已实现", 16, ""),
    (45, "附件编制", "附件07 进度计划", "里程碑/WBS 草稿；与专业软件衔接提示", "plan-workbench schedule", "重要", "已实现", 16, ""),
    (46, "附件编制", "附件08 劳务协议书", "法务类表单与确认项；可生成草稿", "plan-workbench labor_agreement", "重要", "已实现", 12, ""),
    (47, "附件编制", "附件09 保密承诺书", "标准文本锁定；项目/投标人/日期预填", "plan-workbench confidentiality", "重要", "已实现", 10, ""),
    (48, "附件编制", "其他未识别附件占位", "未匹配语义的附件提示工作台稍后实现", "编制区 placeholder", "一般", "已实现", 2, "占位，不阻塞主流程"),
    (49, "投标报价", "报价工作台总览", "三步：材料清单→组价→总价费用", "阶段四；GET .../quotation", "核心", "已实现", 16, ""),
    (50, "投标报价", "上传材料设备清单", "解析 xlsx/xlsm，分开展示有单价/缺单价行", "POST .../material-list", "核心", "已实现", 20, ""),
    (51, "投标报价", "上传当地材料价格表", "写入共享价格库，跨项目可用", "POST .../local-price", "重要", "已实现", 12, ""),
    (52, "投标报价", "历史价自动匹配缺价", "按历史项目材料单价填入缺价行", "POST .../fill-prices", "核心", "已实现", 24, ""),
    (53, "投标报价", "浏览并手工套用历史价", "搜索历史材料，勾选后填入缺价行", "price-history / apply-history-price", "重要", "已实现", 12, ""),
    (54, "投标报价", "从当地价格表补单价", "按共享库匹配填入缺价", "POST .../fill-local-prices", "重要", "已实现", 8, ""),
    (55, "投标报价", "市场行情分析补单价", "LLM 分析市场价填入缺价行", "POST .../fill-market-prices", "重要", "已实现", 16, "依赖 LLM"),
    (56, "投标报价", "供应商邮件询价", "勾选缺价行，SMTP 发询价邮件", "POST .../inquiries", "核心", "已实现", 24, "依赖 SMTP 配置"),
    (57, "投标报价", "收取供应商邮件回复", "拉取邮箱回复并尝试回填单价", "POST .../inquiries/sync-replies", "核心", "已实现", 20, ""),
    (58, "投标报价", "粘贴回复解析单价", "手工粘贴供应商回复文本解析填价", "POST .../inquiries/{id}/reply", "重要", "已实现", 12, ""),
    (59, "投标报价", "生成组价清单", "按模板参数计算综合单价、合价等", "POST .../build-pricing", "核心", "已实现", 28, "需组价模板 xlsx"),
    (60, "投标报价", "预览组价明细", "查看组价行与费率参数", "GET .../preview", "核心", "已实现", 8, ""),
    (61, "投标报价", "导出组价 Excel", "下载填好的组价工作簿", "GET .../export-pricing", "核心", "已实现", 12, ""),
    (62, "投标报价", "总价费用汇总", "在组价基础上汇总总价费用表", "第三步生成总价费用", "重要", "已实现", 16, ""),
    (63, "模板套打", "原附件 Word 套打", "在原招标附件上按字段替换导出，保留版式", "template_fill + export-docx", "核心", "已实现", 40, "含 LLM 补全与下划线策略"),
    (64, "模板套打", "金额中文大写与校验", "承诺函等场景自动/校验大小写金额一致", "附件02 工作台", "重要", "已实现", 8, ""),
    (65, "系统支撑", "LLM 调用配置", "公告/须知/字段/报价等依赖 DashScope", ".env DASHSCOPE_*；health", "核心", "已实现", 4, ""),
    (66, "系统支撑", "邮件 SMTP 配置", "询价发送与回复收取依赖邮件配置", ".env SMTP_*；health", "重要", "已实现", 6, ""),
    (67, "系统支撑", "组价模板文件", "组价导出依赖组价模板 Excel", "templates/组价模板（自用）.xlsx", "核心", "已实现", 4, "部署需随包上传"),
    (68, "系统支撑", "结构化解读表模板", "公告详情/导出依赖解读表模板", "templates/招标文件结构化解读表.xlsx", "重要", "已实现", 4, ""),
]


def allocate_month_hours(rows, total_hours: float = MONTH_HOURS) -> list[float]:
    """按权重把一人一个月工时分摊到各功能，合计精确等于 total_hours。"""
    weights = [float(r[7]) for r in rows]
    wsum = sum(weights) or 1.0
    raw = [total_hours * w / wsum for w in weights]
    hours = [round(x, 1) for x in raw]
    # 修正四舍五入误差到最后一项
    diff = round(total_hours - sum(hours), 1)
    hours[-1] = round(hours[-1] + diff, 1)
    return hours


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "功能清单明细"

    headers = [
        "序号",
        "模块",
        "功能名称",
        "功能说明",
        "入口/页面或API",
        "优先级",
        "状态",
        "分摊工时(小时)",
        "分摊人天",
        "备注",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="微软雅黑", size=10)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    prio_fills = {
        "核心": PatternFill("solid", fgColor="FCE4D6"),
        "重要": PatternFill("solid", fgColor="FFF2CC"),
        "一般": PatternFill("solid", fgColor="E2EFDA"),
    }
    mod_fills = {
        "系统与登录": "D6EAF8",
        "项目管理": "D5F5E3",
        "文件入库": "FCF3CF",
        "招标公告": "F5D0D0",
        "投标须知": "E8DAEF",
        "附件编制": "D6EAF8",
        "投标报价": "FAD7A0",
        "模板套打": "D5F5E3",
        "系统支撑": "D5D8DC",
    }
    hour_fill = PatternFill("solid", fgColor="DEEBF7")

    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    allocated = allocate_month_hours(ROWS, MONTH_HOURS)
    for r, hours in zip(ROWS, allocated):
        days = round(hours / HOURS_PER_DAY, 2)
        ws.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], hours, days, r[8]])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=10):
        prio = row[5].value
        mod = row[1].value
        for i, cell in enumerate(row):
            cell.font = cell_font
            cell.border = thin
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
                horizontal="center" if i in (0, 5, 6, 7, 8) else "left",
            )
        if prio in prio_fills:
            row[5].fill = prio_fills[prio]
        if mod in mod_fills:
            row[1].fill = PatternFill("solid", fgColor=mod_fills[mod])
        row[6].fill = PatternFill("solid", fgColor="C6EFCE")
        row[7].fill = hour_fill
        row[8].fill = hour_fill

    widths = [6, 12, 22, 46, 34, 8, 8, 14, 10, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"

    total_h = round(sum(allocated), 1)
    total_d = round(total_h / HOURS_PER_DAY, 2)
    man_months = round(total_h / MONTH_HOURS, 2)
    sum_row = ws.max_row + 1
    ws.append(
        [
            "",
            "合计",
            "",
            "",
            "",
            "",
            "",
            total_h,
            total_d,
            f"一人一个月总量 {MONTH_HOURS}h（{MONTH_WORKDAYS}天×{HOURS_PER_DAY}h）按复杂度分摊",
        ]
    )
    for c in range(1, 11):
        cell = ws.cell(sum_row, c)
        cell.font = Font(name="微软雅黑", bold=True, size=10)
        cell.border = thin
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 模块汇总
    ws2 = wb.create_sheet("模块汇总")
    ws2.append(["模块", "功能数", "核心", "重要", "一般", "分摊工时(小时)", "分摊人天", "占一个月比例"])
    stats: dict = defaultdict(lambda: Counter())
    hours_by_mod: dict = defaultdict(float)
    for r, h in zip(ROWS, allocated):
        stats[r[1]][r[5]] += 1
        stats[r[1]]["all"] += 1
        hours_by_mod[r[1]] += h
    order = [
        "系统与登录",
        "项目管理",
        "文件入库",
        "招标公告",
        "投标须知",
        "附件编制",
        "投标报价",
        "模板套打",
        "系统支撑",
    ]
    for m in order:
        c = stats[m]
        h = round(hours_by_mod[m], 1)
        ws2.append(
            [
                m,
                c["all"],
                c["核心"],
                c["重要"],
                c["一般"],
                h,
                round(h / HOURS_PER_DAY, 2),
                f"{round(h / MONTH_HOURS * 100, 1)}%",
            ]
        )
    ws2.append(
        [
            "合计",
            len(ROWS),
            sum(1 for r in ROWS if r[5] == "核心"),
            sum(1 for r in ROWS if r[5] == "重要"),
            sum(1 for r in ROWS if r[5] == "一般"),
            total_h,
            total_d,
            "100%",
        ]
    )
    for c in range(1, 9):
        cell = ws2.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, w in enumerate([14, 10, 8, 8, 8, 14, 10, 12], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=8):
        for cell in row:
            cell.font = cell_font
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws3 = wb.create_sheet("说明")
    notes = [
        ["项目", "AI 投标报价系统（aivl）功能清单明细"],
        ["生成依据", "当前代码实现（后端 FastAPI + 前端 Vue）"],
        ["业务主链路", "登录 → 项目 → 上传 ZIP → 招标公告分析 → 投标须知（附件编制） → 投标报价"],
        ["优先级说明", "核心=主路径必做；重要=常用增强；一般=辅助/运维/占位"],
        ["工时口径", f"一人一个月：双休、每天 {HOURS_PER_DAY} 小时、约 {MONTH_WORKDAYS} 个工作日 = {MONTH_HOURS} 小时"],
        ["分摊方式", "按各功能复杂度权重等比分摊，全部功能工时合计精确等于一人一个月"],
        ["人天换算", f"1 人天 = {HOURS_PER_DAY} 小时；合计 {total_h} 小时 = {total_d} 人天 = {man_months} 人月"],
        ["部署注意", "线上需配置 DASHSCOPE_API_KEY、SMTP_*，并放置 templates 下组价/解读表模板"],
    ]
    for r in notes:
        ws3.append(r)
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 96
    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, min_col=1, max_col=2):
        row[0].font = Font(name="微软雅黑", bold=True, size=10)
        row[1].font = Font(name="微软雅黑", size=10)

    try:
        wb.save(OUT)
        saved = OUT
    except PermissionError:
        wb.save(OUT_ALT)
        saved = OUT_ALT
    print(saved)
    print("rows", len(ROWS))
    print(f"total_hours={total_h} person_days={total_d} man_months={man_months}")
    print("sample:", list(zip([r[2] for r in ROWS[:3]], allocated[:3])))


if __name__ == "__main__":
    main()
