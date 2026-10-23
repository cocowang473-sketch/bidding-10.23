# -*- coding: utf-8 -*-
"""生成面向客户报价的「智能化投标决策与报价协同平台」能力说明（商务美化版）。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "release" / "智能化投标决策与报价协同平台-客户报价能力说明.xlsx"
OUT_ALT = ROOT / "release" / "智能化投标决策与报价协同平台-客户报价能力说明-V1.2.xlsx"

MODULE_COLORS = {
    "智能招标文件认知引擎": "1F4E79",
    "历史数据智能分析中心": "2E75B6",
    "合规与附件智能编制中枢": "548235",
    "智能组价与成本推演引擎": "C65911",
    "供应链询价协同与邮件编排": "833C0C",
    "决策看板与交付编排": "7030A0",
    "安全、权限与交付保障": "595959",
}

# 序号, 能力域, 能力项, 客户可见价值说明, 技术复杂度标签, 交付阶段, 备注
CAPABILITIES = [
    (
        1,
        "智能招标文件认知引擎",
        "多模态招标包结构化解析",
        "支持完整招标数据包一键接入，自动完成文档体系建模、章节语义切分与关键实体抽取，形成结构化招标信息对象，显著缩短人工翻阅与摘录周期。",
        "高",
        "一期",
        "",
    ),
    (
        2,
        "智能招标文件认知引擎",
        "公告/须知多维合规要素抽取",
        "对招标公告、投标须知进行深度语义解析，沉淀资格门槛、商务条款、技术约束、时间节点与风险提示，输出结构化决策要素清单，支撑后续编制与复核。",
        "高",
        "一期",
        "",
    ),
    (
        3,
        "智能招标文件认知引擎",
        "跨文档一致性与冲突识别",
        "对项目名称、金额口径、人员资质、进度节点等关键信息进行跨附件联动校验，提前暴露前后矛盾与遗漏项，降低废标与返工风险。",
        "高",
        "一期",
        "",
    ),
    (
        4,
        "智能招标文件认知引擎",
        "法语标书内容识别与语义理解",
        "面向法语（及中法混排）招标文件，自动识别章节结构、工程描述、规格型号与商务条款等关键内容，完成跨语言语义理解与结构化抽取，降低人工翻译与摘录成本。",
        "很高",
        "一期",
        "",
    ),
    (
        5,
        "智能招标文件认知引擎",
        "法语工程量/材料清单解析入模",
        "支持法语版工程量清单、材料设备表等表格类文件的字段识别与入模，处理 designation、unité、quantité 等典型表达，并与组价模型字段对齐。",
        "很高",
        "一期",
        "",
    ),
    (
        6,
        "智能招标文件认知引擎",
        "中法术语对齐与双语对照输出",
        "在解析与组价过程中建立中法术语对照，支持关键描述双语呈现与复核，便于国内编制团队与海外商务协同作业。",
        "高",
        "一期",
        "",
    ),
    (
        7,
        "历史数据智能分析中心",
        "企业历史报价与采购数据沉淀",
        "将历史项目报价结果、材料/机械采购实绩、税率与币种信息等持续沉淀为可分析的历史数据集，为后续智能回填提供高质量样本基础。",
        "高",
        "一期",
        "",
    ),
    (
        8,
        "历史数据智能分析中心",
        "AI 历史相似项识别与关联分析",
        "基于大模型对工程描述、规格参数、单位体系进行语义理解（含法语描述场景），在历史数据中识别同类材料与相近规格，建立可复核的关联关系，支撑单价与税率建议。",
        "很高",
        "一期",
        "",
    ),
    (
        9,
        "历史数据智能分析中心",
        "单价智能回填",
        "对工程量清单中缺价或待核价条目，自动引用历史分析得到的参考单价完成主材费等字段回填；支持未匹配筛选、勾选重算与人工确认，形成人机协同闭环。",
        "很高",
        "一期",
        "",
    ),
    (
        10,
        "历史数据智能分析中心",
        "税率/币种与价税口径推算",
        "结合历史采购样本中的币种、税率及含税/不含税口径，自动完成税率识别、价税换算与口径对齐，减少手工换算差错，保障组价结果可比、可审。",
        "高",
        "一期",
        "",
    ),
    (
        11,
        "历史数据智能分析中心",
        "历史数据运营与权限治理",
        "支持管理员对历史价格与采购样本进行受控入库、覆盖更新与留痕管理；业务侧消费分析结果，核心成本数据与日常作业权限隔离。",
        "中高",
        "一期",
        "",
    ),
    (
        12,
        "合规与附件智能编制中枢",
        "动态附件清单识别与流程编排",
        "自动识别须知要求的附件体系，按条件分支驱动“是否纳入/无需提交”，并以向导式流程组织编制、校验与完结，适配不同招标文件差异。",
        "高",
        "一期",
        "",
    ),
    (
        13,
        "合规与附件智能编制中枢",
        "模板语义填空与主数据贯通",
        "从原版 Word 模板识别待填语义槽位，联动企业主数据与项目基础信息自动预填；支持资料导入辅助回填，减少重复录入。",
        "很高",
        "一期",
        "",
    ),
    (
        14,
        "合规与附件智能编制中枢",
        "原版式合规套打导出",
        "在保留招标方原附件版式与签章留白的前提下完成字段回写导出，满足正式投标文件观感与格式合规要求。",
        "很高",
        "一期",
        "",
    ),
    (
        15,
        "合规与附件智能编制中枢",
        "人员组织与计划类附件协同",
        "覆盖授权委托、承诺函、人员履历、组织架构、劳动力/设备/进度计划及保密协议等关键附件的结构化维护与草稿生成。",
        "高",
        "一期",
        "",
    ),
    (
        16,
        "智能组价与成本推演引擎",
        "工程量清单智能入模",
        "将工程量清单自动映射至组价模型字段体系，完成描述规范化、单位对齐与明细行建模，为定额匹配与主材费推演奠定统一数据底座。",
        "高",
        "一期",
        "",
    ),
    (
        17,
        "智能组价与成本推演引擎",
        "定额含量智能匹配",
        "基于历史定额样本与 AI 语义理解，对人工/辅材/机械等定额含量进行匹配回填，支持批量重算与结果复核，提升组价效率与口径一致性。",
        "高",
        "一期",
        "",
    ),
    (
        18,
        "智能组价与成本推演引擎",
        "主材费与综合单价联动推演",
        "在 AI 历史单价回填与税率口径对齐基础上，按企业组价模型完成主材费、综合单价、合价及总价费用汇总推演，并支持成果导出。",
        "很高",
        "一期",
        "",
    ),
    (
        19,
        "智能组价与成本推演引擎",
        "参数化费率与场景重算",
        "支持人工费率、运杂、综合取费等关键参数配置与场景重算，便于商务在不同策略假设下快速产出可比报价方案。",
        "中高",
        "一期",
        "",
    ),
    (
        20,
        "供应链询价协同与邮件编排",
        "缺口材料智能询价编排",
        "对历史分析仍无法可靠定价的材料项发起结构化询价任务，自动生成询价清单与邮件正文，支持批量定向发送至合格供方。",
        "高",
        "一期",
        "",
    ),
    (
        21,
        "供应链询价协同与邮件编排",
        "邮件通道收发与会话归集",
        "通过企业邮箱通道完成询价发出与回复收取，按项目/询价批次归集会话上下文，避免报价信息散落在个人邮箱中。",
        "高",
        "一期",
        "",
    ),
    (
        22,
        "供应链询价协同与邮件编排",
        "回复语义解析与价税回写",
        "对供应商回复（自然语言、表格混排、中英文/法文混杂）进行 AI 解析，抽取单价、币种、税率等要素并回写报价模型，缩短询价闭环周期。",
        "很高",
        "一期",
        "",
    ),
    (
        23,
        "供应链询价协同与邮件编排",
        "历史价与回函价比对决策支持",
        "在历史分析建议价与供应商回函价之间提供比对视图，辅助商务完成最终单价与税率口径确认。",
        "高",
        "二期增强",
        "",
    ),
    (
        24,
        "决策看板与交付编排",
        "项目分阶段作战视图",
        "以“资料接入→公告认知→须知合规→报价推演”阶段视图管理项目进度，关键节点状态可视化，便于项目经理与商务协同推进。",
        "中高",
        "一期",
        "",
    ),
    (
        25,
        "决策看板与交付编排",
        "可追溯依据链",
        "关键单价、税率与组价结果保留来源线索（历史数据分析、邮件回函、人工确认等），满足内部审计与报价复核需要。",
        "中高",
        "一期",
        "",
    ),
    (
        26,
        "决策看板与交付编排",
        "标准化成果包输出",
        "一键输出结构化解读、套打附件、组价明细与总价汇总等成果包，支撑对内评审与对外投标文件组装。",
        "中高",
        "一期",
        "",
    ),
    (
        27,
        "安全、权限与交付保障",
        "身份认证与会话治理",
        "企业账号登录、会话校验与过期重认证，保障业务操作可审计、可回收。",
        "中",
        "一期",
        "",
    ),
    (
        28,
        "安全、权限与交付保障",
        "管理域与业务域隔离",
        "历史价格样本维护与日常投标作业权限分离，降低敏感成本数据扩散风险。",
        "中高",
        "一期",
        "",
    ),
    (
        29,
        "安全、权限与交付保障",
        "私有化/专属环境交付适配",
        "支持按客户环境完成模型服务、邮件通道与历史数据存储的专属部署与联调，满足数据不出域等合规诉求（以合同约定范围为准）。",
        "高",
        "交付项",
        "",
    ),
]

OVERVIEW = [
    ("方案定位", "面向机电/工程类（含海外法语标）投标场景的智能化决策与报价协同平台，覆盖招标文件认知、法语标书识别、合规附件编制、历史数据驱动的智能组价，以及供应链询价闭环。"),
    ("核心差异", "具备法语及中法混排标书内容识别能力；并以大模型分析企业历史报价与采购数据，自动完成相似材料识别、单价回填、税率与价税口径推算；结合邮件编排完成供方报价采集。"),
    ("客户收益", "缩短法语/中文标书研读与组价周期；降低漏项与价税口径不一致风险；提升价格依据可追溯性；让历史项目经验持续复利到新报价中。"),
    ("交付形态", "B/S 工作台 + 多语言标书认知 + 历史数据分析 + 邮件协同通道；支持分阶段上线与按项目试用。"),
]

ARCHITECTURE = [
    ("L1 接入层", "招标数据包接入、工程量清单/历史采购数据导入、邮件通道接入"),
    ("L2 认知层", "中/法语文档解析、语义抽取、跨文档校验、模板槽位识别、中法术语对齐"),
    ("L3 分析层", "历史数据治理、AI 相似项识别、单价回填、税率与价税口径推算"),
    ("L4 推演层", "定额匹配、主材费联动、组价参数推演、总价汇总"),
    ("L5 协同层", "询价任务编排、邮件收发、回复解析回写、人工确认闭环"),
    ("L6 治理层", "权限隔离、操作留痕、成果导出、部署与运维保障"),
]

PHASES = [
    ("一期（标准交付）", "中法标书认知、须知与附件编制、历史数据智能分析、单价回填与税率推算、组价推演、邮件询价收发闭环、权限与成果导出"),
    ("二期（增强建议）", "历史价与回函价融合决策看板、供应商评分、跨项目组合分析、报价策略沙盘、更多语种扩展"),
    ("可选增值", "专属模型适配、行业中法术语库共建、与客户 ERP/成本系统对接、历史数据清洗与冷启动辅导"),
]


def style_sheet(ws):
    return Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    cover = wb.active
    cover.title = "01-方案总览"
    cover["A1"] = "智能化投标决策与报价协同平台"
    cover["A2"] = "客户报价能力说明（商务版）"
    cover["A3"] = "版本：V1.2｜用途：商务对外报价与方案沟通｜不含内部实现细节｜含法语标书识别能力"
    cover["A1"].font = Font(name="微软雅黑", size=20, bold=True, color="1F4E79")
    cover["A2"].font = Font(name="微软雅黑", size=14, bold=True, color="2E75B6")
    cover["A3"].font = Font(name="微软雅黑", size=10, color="666666")

    cover["A5"] = "一、方案摘要"
    cover["A5"].font = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
    cover["A5"].fill = PatternFill("solid", fgColor="1F4E79")
    cover.merge_cells("A5:B5")

    cover["A6"], cover["B6"] = "维度", "说明"
    for col in ("A", "B"):
        cell = cover[f"{col}6"]
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (k, v) in enumerate(OVERVIEW, start=7):
        cover[f"A{i}"] = k
        cover[f"B{i}"] = v
        cover[f"A{i}"].font = Font(name="微软雅黑", bold=True)
        cover[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        cover[f"A{i}"].alignment = Alignment(vertical="top")
        cover.row_dimensions[i].height = 48

    cover["A12"] = "二、能力架构（对客表述）"
    cover["A12"].font = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
    cover["A12"].fill = PatternFill("solid", fgColor="1F4E79")
    cover.merge_cells("A12:B12")
    cover["A13"] = "层级"
    cover["B13"] = "能力构成"
    for col in ("A", "B"):
        cover[f"{col}13"].font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cover[f"{col}13"].fill = PatternFill("solid", fgColor="2E75B6")

    for i, (k, v) in enumerate(ARCHITECTURE, start=14):
        cover[f"A{i}"] = k
        cover[f"B{i}"] = v
        cover[f"B{i}"].alignment = Alignment(wrap_text=True)
        cover.row_dimensions[i].height = 30

    cover["A21"] = "三、交付节奏建议"
    cover["A21"].font = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
    cover["A21"].fill = PatternFill("solid", fgColor="1F4E79")
    cover.merge_cells("A21:B21")
    cover["A22"] = "阶段"
    cover["B22"] = "范围要点"
    for col in ("A", "B"):
        cover[f"{col}22"].font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cover[f"{col}22"].fill = PatternFill("solid", fgColor="2E75B6")
    for i, (k, v) in enumerate(PHASES, start=23):
        cover[f"A{i}"] = k
        cover[f"B{i}"] = v
        cover[f"B{i}"].alignment = Alignment(wrap_text=True)
        cover.row_dimensions[i].height = 40

    cover["A27"] = (
        "声明：本说明用于商务沟通与报价范围对齐，描述的是可交付业务能力与专业复杂度；"
        "具体接口、模型配置、算法细节与内部工具链不在对外材料披露范围。最终以合同与 SOW 为准。"
    )
    cover["A27"].font = Font(name="微软雅黑", size=9, italic=True, color="888888")
    cover.merge_cells("A27:B27")
    cover["A27"].alignment = Alignment(wrap_text=True)
    cover.row_dimensions[27].height = 40

    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 88

    ws = wb.create_sheet("02-能力清单明细")
    thin = style_sheet(ws)
    headers2 = ["序号", "能力域", "能力项", "客户价值与专业内涵", "复杂度", "建议阶段"]
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    for c, h in enumerate(headers2, start=1):
        cell = ws.cell(1, c, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(CAPABILITIES) + 1}"

    for r_idx, row in enumerate(CAPABILITIES, start=2):
        seq, domain, name, desc, complexity, phase, _note = row
        values = [seq, domain, name, desc, complexity, phase]
        domain_color = MODULE_COLORS.get(domain, "2E75B6")
        for c, val in enumerate(values, start=1):
            cell = ws.cell(r_idx, c, val)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(name="微软雅黑", size=10)
            if c == 2:
                cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=domain_color)
            if c == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if complexity in ("很高", "高"):
                    cell.font = Font(name="微软雅黑", size=10, bold=True, color="C00000")
        ws.row_dimensions[r_idx].height = 56

    for i, w in enumerate([6, 28, 26, 72, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    talk = wb.create_sheet("03-商务沟通要点")
    talk["A1"] = "对外沟通建议（内部参考）"
    talk["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    tips = [
        "强调“AI 分析历史报价/采购数据做单价回填与税率推算”，避免说成“Excel 对一下表”。",
        "强调“法语及中法混排标书识别 + 中法术语对齐”，体现海外项目专业门槛，避免只说“能翻译”。",
        "强调“跨文档合规认知与冲突识别”，避免展开具体提示词或字段映射规则。",
        "强调“询价邮件编排与回复语义回写闭环”，体现供应链协同专业度。",
        "复杂度用语可用：跨语言语义理解、历史关联分析、价税口径推算、人机协同确认、可追溯依据链。",
        "不要对外展示：内部 API 路径、模型厂商名、表头别名规则、简单字符串匹配等实现细节。",
        "报价时可按能力域打包：认知引擎（含法语） / 历史数据分析 / 编制中枢 / 组价引擎 / 询价协同 / 治理交付。",
        "若客户追问算法，统一答复为“基于企业历史数据与多语言语义理解的大模型分析，细节属实施交付物”。",
    ]
    talk["A2"] = "建议话术与边界"
    talk["A2"].font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    talk["A2"].fill = PatternFill("solid", fgColor="2E75B6")
    talk.merge_cells("A2:B2")
    talk["A3"] = "序号"
    talk["B3"] = "要点"
    for col in ("A", "B"):
        talk[f"{col}3"].font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        talk[f"{col}3"].fill = PatternFill("solid", fgColor="5B9BD5")
    for i, t in enumerate(tips, start=4):
        talk[f"A{i}"] = i - 3
        talk[f"B{i}"] = t
        talk[f"B{i}"].alignment = Alignment(wrap_text=True)
        talk.row_dimensions[i].height = 28
    talk.column_dimensions["A"].width = 8
    talk.column_dimensions["B"].width = 96

    pack = wb.create_sheet("04-报价模块包装参考")
    pack["A1"] = "建议对外报价模块（可按包报价）"
    pack["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    pack_headers = ["模块包", "包含能力摘要", "专业卖点", "建议计价口径"]
    for c, h in enumerate(pack_headers, start=1):
        cell = pack.cell(2, c, h)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

    packs = [
        (
            "A. 招标认知与合规包（含法语）",
            "多模态解析、法语标书识别、中法术语对齐、公告/须知抽取、跨文档校验、附件套打",
            "打通海外法语标读懂与国内编制协同",
            "按项目数 / 年订阅",
        ),
        (
            "B. 历史数据智能分析包",
            "历史数据沉淀、AI 相似项识别、单价回填、税率/币种与价税口径推算",
            "用 AI 把历史经验变成可复用的定价能力",
            "建设费 + 数据运营年费",
        ),
        (
            "C. 智能组价推演包",
            "工程量入模、定额匹配、主材费联动、综合单价与总价推演导出",
            "人机协同的成本推演，而非手工拼表",
            "按项目数 / 并发席位",
        ),
        (
            "D. 供应链询价协同包",
            "询价编排、邮件收发归集、回复语义解析与价税回写、历史价比对",
            "打通“缺价→询价→回函→入模”闭环",
            "按席位 + 邮件通道实施",
        ),
        (
            "E. 专属交付与治理包",
            "权限隔离、私有化适配、联调上线、培训与保驾",
            "满足数据安全与落地可用性",
            "一次性实施 + 运维",
        ),
    ]
    for r, row in enumerate(packs, start=3):
        for c, val in enumerate(row, start=1):
            cell = pack.cell(r, c, val)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = thin
        pack.row_dimensions[r].height = 48
    for i, w in enumerate([26, 48, 36, 22], start=1):
        pack.column_dimensions[get_column_letter(i)].width = w

    try:
        wb.save(OUT)
        print(f"written: {OUT}")
    except PermissionError:
        wb.save(OUT_ALT)
        print(f"原文件被占用，已写入: {OUT_ALT}")


if __name__ == "__main__":
    main()
