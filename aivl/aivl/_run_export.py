# -*- coding: utf-8 -*-
import json
import time
from collections import Counter
from pathlib import Path

from app.excel_export import build_workbook_bytes, pick_primary_docs, _infer_focus_keys, _score_doc

pid = 2
print("=== primary ===", flush=True)
for d in pick_primary_docs(pid, limit=8):
    print(_score_doc(d), _infer_focus_keys(d["filename"], d.get("rel_path") or ""), d["filename"][:90], flush=True)

print("\n=== building workbook (LLM) ===", flush=True)
t0 = time.time()
content, meta = build_workbook_bytes(pid)
out = Path("scratch") / "export_verify.xlsx"
out.parent.mkdir(exist_ok=True)
out.write_bytes(content)
print("wrote", out, "bytes", len(content), f"in {time.time()-t0:.1f}s", flush=True)
ext = meta.get("extract") or {}
print("row_count", meta.get("row_count"), "chunks", ext.get("chunks_done"), flush=True)
print("errors", ext.get("llm_errors") or meta.get("llm_error"), flush=True)
print("fallbacks", json.dumps(ext.get("fallbacks") or [], ensure_ascii=False), flush=True)
for p in ext.get("primary_docs") or []:
    print(" doc", p, flush=True)

from openpyxl import load_workbook

wb = load_workbook(out, data_only=True)
ws = wb["01_招标解读总表"]
keys = Counter()
src = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[1]:
        continue
    keys[str(row[1]).split()[0]] += 1
    if row[7]:
        src[str(row[7])[:90]] += 1
print("KEY counts", dict(keys), "total", sum(keys.values()), flush=True)
print("sample expect ~242; K06~100 K03~48", flush=True)
print("TOP sources:", flush=True)
for s, c in src.most_common(12):
    print(c, s, flush=True)
