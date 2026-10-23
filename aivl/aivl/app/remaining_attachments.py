"""附件05–09：劳动力/设备/进度/劳务协议/保密承诺 工作台。"""
from __future__ import annotations

import io
import re
import uuid
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .instructions_analysis import (
    STATUS_LABEL,
    get_or_init_instructions,
    update_instructions,
)
from .text_extract import extract_text
from .ingest import absolute_doc_path

KIND_HANDLERS = ("labor_plan", "equipment_plan", "schedule", "labor_agreement", "confidentiality")


def _find(state: Dict[str, Any], *, code: Optional[str] = None, kind: Optional[str] = None):
    for a in state.get("attachments") or []:
        if code and a.get("code") == code:
            return a
        if kind and a.get("kind") == kind:
            return a
    return None


def _uid() -> str:
    return str(uuid.uuid4())


def _drafting(status: Optional[str]) -> str:
    if status in ("to_prepare", "pending_decision", "pending_p6"):
        return "drafting"
    return status or "drafting"


# ---------- defaults ----------

DEFAULT_PHASES = [
    "深化设计",
    "材料进场",
    "安装施工",
    "调试验收",
]

DEFAULT_TRADES = ["电工", "水暖工", "通风工", "焊工", "普工"]

DEFAULT_EQUIPMENT = [
    {"name": "电焊机", "spec": "ZX7-400", "qty": "4"},
    {"name": "套丝机", "spec": "2-4寸", "qty": "2"},
    {"name": "电动试压泵", "spec": "常用", "qty": "2"},
    {"name": "高空作业平台", "spec": "自行式", "qty": "1"},
    {"name": "周转脚手架", "spec": "门式", "qty": "若干"},
]

DEFAULT_MILESTONES = [
    {"name": "深化设计完成", "type": "深化设计"},
    {"name": "主要材料到场", "type": "材料进场"},
    {"name": "机电安装开工", "type": "施工组织"},
    {"name": "系统调试完成", "type": "施工组织"},
    {"name": "竣工验收", "type": "里程碑"},
]


def _new_phase(seq: int = 1, **kw) -> Dict[str, Any]:
    return {
        "id": kw.get("id") or _uid(),
        "seq": int(kw.get("seq") or seq),
        "phase_name": kw.get("phase_name") or "",
        "trade": kw.get("trade") or "",
        "level": kw.get("level") or "",
        "headcount": kw.get("headcount") if kw.get("headcount") is not None else "",
        "notes": kw.get("notes") or "",
    }


def _new_curve(period: str, headcount: Any = 0, peak: bool = False) -> Dict[str, Any]:
    return {"id": _uid(), "period": period, "headcount": headcount, "peak": peak}


def _new_equip(seq: int = 1, **kw) -> Dict[str, Any]:
    return {
        "id": kw.get("id") or _uid(),
        "seq": int(kw.get("seq") or seq),
        "name": kw.get("name") or "",
        "spec": kw.get("spec") or "",
        "age_years": kw.get("age_years") or "",
        "qty": kw.get("qty") or "",
        "location": kw.get("location") or "",
        "mobilize_date": kw.get("mobilize_date") or "",
        "source": kw.get("source") or "",
        "source_note": kw.get("source_note") or "",
        "qty_note": kw.get("qty_note") or "",
        "ready_now": kw.get("ready_now") if "ready_now" in kw else None,
    }


def _new_ms(seq: int = 1, **kw) -> Dict[str, Any]:
    return {
        "id": kw.get("id") or _uid(),
        "seq": int(kw.get("seq") or seq),
        "name": kw.get("name") or "",
        "type": kw.get("type") or "",
        "planned_start": kw.get("planned_start") or "",
        "planned_finish": kw.get("planned_finish") or "",
        "notes": kw.get("notes") or "",
    }


def ensure_labor_plan(form: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(form or {})
    phases = out.get("phases")
    if not isinstance(phases, list) or not phases:
        phases = []
        seq = 1
        for ph in DEFAULT_PHASES:
            for trade in DEFAULT_TRADES[:3]:
                phases.append(_new_phase(seq, phase_name=ph, trade=trade, level="中级", headcount=""))
                seq += 1
        out["phases"] = phases
    else:
        out["phases"] = [_new_phase(i + 1, **{k: v for k, v in (p if isinstance(p, dict) else {}).items() if k != "seq"}) for i, p in enumerate(phases)]
    if not isinstance(out.get("curve_points"), list) or not out["curve_points"]:
        out["curve_points"] = _curve_from_phases(out["phases"])
    out.setdefault("peak_headcount", "")
    out.setdefault(
        "chinese_source",
        {
            "hq_dispatch": {"eligible": None, "count": "", "note": ""},
            "from_other_projects": {"passport_done": None, "visa_done": None, "count": "", "note": ""},
            "trained_certified": {"count": "", "note": ""},
        },
    )
    out.setdefault("credential_notes", "")
    out.setdefault("voucher_attached", False)
    out.setdefault("notes", "")
    return out


def _curve_from_phases(phases: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    totals: Dict[str, int] = {}
    for p in phases:
        name = str(p.get("phase_name") or "阶段").strip() or "阶段"
        try:
            n = int(float(str(p.get("headcount") or "0").replace(",", "") or 0))
        except Exception:
            n = 0
        totals[name] = totals.get(name, 0) + n
    if not totals:
        return [_new_curve(ph, 0) for ph in DEFAULT_PHASES]
    peak = max(totals.values()) if totals else 0
    return [_new_curve(k, v, peak=(v == peak and v > 0)) for k, v in totals.items()]


def ensure_equipment_plan(form: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(form or {})
    rows = out.get("equipment_rows")
    if not isinstance(rows, list) or not rows:
        out["equipment_rows"] = [_new_equip(i + 1, **e) for i, e in enumerate(DEFAULT_EQUIPMENT)]
    else:
        out["equipment_rows"] = [
            _new_equip(i + 1, **{k: v for k, v in (r if isinstance(r, dict) else {}).items() if k != "seq"})
            for i, r in enumerate(rows)
        ]
    out.setdefault(
        "confirm_flags",
        {
            "source_explained": False,
            "mobilize_aligned_schedule": False,
            "voucher_prepared": False,
            "ready_confirmed": False,
        },
    )
    out.setdefault("voucher_note", "")
    out.setdefault("notes", "")
    return out


def ensure_schedule(form: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(form or {})
    ms = out.get("wbs_milestones")
    if not isinstance(ms, list) or not ms:
        out["wbs_milestones"] = [_new_ms(i + 1, **m) for i, m in enumerate(DEFAULT_MILESTONES)]
    else:
        out["wbs_milestones"] = [
            _new_ms(i + 1, **{k: v for k, v in (m if isinstance(m, dict) else {}).items() if k != "seq"})
            for i, m in enumerate(ms)
        ]
    out.setdefault(
        "scope_covers",
        {"detail_design": False, "material_delivery": False, "construction_org": False},
    )
    out.setdefault(
        "p6",
        {"export_status": "pending", "import_status": "none", "file_name": "", "last_synced_at": None},
    )
    out.setdefault("engineer_notes", "")
    out.setdefault("confirmed", False)
    return out


LABOR_AGREE_FIELDS = [
    {"key": "party_a", "label": "甲方", "source": "project", "required": True, "hint": "招标人/甲方", "evidence": "甲方："},
    {"key": "party_b", "label": "乙方（投标人）", "source": "enterprise", "required": True, "hint": "投标人全称", "evidence": "乙方："},
    {"key": "uscc", "label": "统一社会信用代码", "source": "enterprise", "required": False, "hint": "信用代码", "evidence": "统一社会信用代码："},
    {"key": "legal_rep", "label": "法定代表人", "source": "legal_rep", "required": True, "hint": "法人姓名", "evidence": "法定代表人："},
    {"key": "authorized_agent", "label": "授权委托人", "source": "agent", "required": False, "hint": "授权委托人", "evidence": "授权委托人："},
    {"key": "sign_date", "label": "签署日期", "source": "manual", "required": True, "hint": "YYYY-MM-DD", "evidence": "日期："},
]

CONF_FIELDS = [
    {"key": "employer", "label": "致（招标人）", "source": "project", "required": True, "hint": "招标人", "evidence": "致："},
    {"key": "project_name", "label": "项目名称", "source": "project", "required": True, "hint": "项目全称", "evidence": "有关 项目的信息"},
    {"key": "bidder_name", "label": "投标人", "source": "enterprise", "required": True, "hint": "投标人全称", "evidence": "投 标 人："},
    {"key": "legal_rep", "label": "法定代表人", "source": "legal_rep", "required": True, "hint": "法人", "evidence": "法定单标人："},
    {"key": "sign_date", "label": "签署日期", "source": "manual", "required": True, "hint": "YYYY-MM-DD", "evidence": "日 期："},
]


def ensure_labor_agreement(form: Dict[str, Any], base: Dict[str, Any], master: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(form or {})
    fields = out.get("fields")
    if not isinstance(fields, list) or not fields:
        fields = []
        for f in LABOR_AGREE_FIELDS:
            item = dict(f)
            item["value"] = _prefill_value(f["key"], base, master)
            fields.append(item)
        out["fields"] = fields
    out.setdefault("locked_body_preview", "")
    out.setdefault(
        "legal_checklist",
        {
            "use_employer_quota": None,
            "non_employer_work": None,
            "accept_advance_deduct": None,
            "need_liability_boundary": None,
            "legal_reviewed": False,
            "legal_reviewer": "",
            "legal_notes": "",
        },
    )
    out.setdefault("sign_reminder", ["甲乙双方签字盖章", "法务审查责任边界与扣款条款"])
    return out


def ensure_confidentiality(form: Dict[str, Any], base: Dict[str, Any], master: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(form or {})
    fields = out.get("fields")
    if not isinstance(fields, list) or not fields:
        fields = []
        for f in CONF_FIELDS:
            item = dict(f)
            item["value"] = _prefill_value(f["key"], base, master)
            fields.append(item)
        out["fields"] = fields
    out.setdefault("locked_body_preview", "")
    out.setdefault("scope_confirm", False)
    out.setdefault("duration_until_acceptance", False)
    out.setdefault("supplier_nda_confirm", False)
    out.setdefault("notes", "")
    out.setdefault("sign_reminder", ["投标人签章", "确认不对供应商泄密直至竣工验收"])
    return out


def _prefill_value(key: str, base: Dict[str, Any], master: Dict[str, Any]) -> str:
    ent = (master or {}).get("enterprise") or {}
    legal = (master or {}).get("legal_rep") or {}
    agent = (master or {}).get("agent") or {}
    proj = (master or {}).get("project") or {}
    mapping = {
        "party_a": base.get("employer") or proj.get("employer_name") or "中建阿尔及利亚股份有限公司",
        "party_b": base.get("bidder_name") or ent.get("bidder_name") or "",
        "uscc": ent.get("credit_code") or "",
        "legal_rep": base.get("legal_representative") or legal.get("legal_rep_name") or "",
        "authorized_agent": base.get("agent_name") or agent.get("agent_name") or "",
        "sign_date": date.today().isoformat(),
        "employer": base.get("employer") or proj.get("employer_name") or "中建阿尔及利亚股份有限公司",
        "project_name": base.get("project_name") or proj.get("project_name") or "",
        "bidder_name": base.get("bidder_name") or ent.get("bidder_name") or "",
    }
    return str(mapping.get(key) or "")


def _org_master(project_id: int) -> Dict[str, Any]:
    from .attachment_workbench import get_org_master

    return get_org_master(project_id)


def _extract_body(project_id: int, att: Dict[str, Any]) -> str:
    path = absolute_doc_path(project_id, att["rel_path"])
    text, _ = extract_text(path, max_chars=12000)
    # 粗清洗
    lines = []
    for line in (text or "").replace("\r", "\n").split("\n"):
        cn = sum(1 for ch in line if "\u4e00" <= ch <= "\u9fff")
        if cn >= 2:
            lines.append(line.strip())
    return "\n".join(lines)


def ensure_form_for_kind(project_id: int, att: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
    kind = att.get("kind")
    form = dict(att.get("form_data") or {})
    base = state.get("base_data") or {}
    master = _org_master(project_id)
    if kind == "labor_plan":
        return ensure_labor_plan(form)
    if kind == "equipment_plan":
        return ensure_equipment_plan(form)
    if kind == "schedule":
        return ensure_schedule(form)
    if kind == "labor_agreement":
        form = ensure_labor_agreement(form, base, master)
        if not form.get("locked_body_preview"):
            form["locked_body_preview"] = _extract_body(project_id, att)[:4000]
        # 同步 conditions.labor_checks
        checks = (state.get("conditions") or {}).get("labor_checks") or {}
        lc = form["legal_checklist"]
        for k in ("use_employer_quota", "non_employer_work", "accept_advance_deduct", "need_liability_boundary"):
            if lc.get(k) is None and k in checks:
                lc[k] = checks.get(k)
        return form
    if kind == "confidentiality":
        form = ensure_confidentiality(form, base, master)
        if not form.get("locked_body_preview"):
            form["locked_body_preview"] = _extract_body(project_id, att)[:4000]
        return form
    raise ValueError(f"不支持的附件类型: {kind}")


def get_plan_workbench(project_id: int, code: str) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find(state, code=code)
    if not att or att.get("kind") not in KIND_HANDLERS:
        raise ValueError("当前附件不在附件05–09工作台范围")
    form = ensure_form_for_kind(project_id, att, state)
    # 持久化默认结构
    status = att.get("status")
    if status in ("to_prepare", "pending_decision") and att.get("kind") != "schedule":
        status = "drafting"
    if att.get("kind") == "schedule" and status == "to_prepare":
        status = "pending_p6"
    update_instructions(project_id, {"attachment": {"code": code, "form_data": form, "status": status}})
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find(state, code=code)
    form = ensure_form_for_kind(project_id, att, state)
    sch = _find(state, kind="schedule")
    extra = {
        "curve_svg": render_labor_curve_svg(form.get("curve_points") or []) if att.get("kind") == "labor_plan" else None,
        "schedule_ref": {
            "code": (sch or {}).get("code"),
            "status": (sch or {}).get("status"),
            "status_label": STATUS_LABEL.get((sch or {}).get("status") or "", (sch or {}).get("status")),
        },
        "conditions": state.get("conditions") or {},
        "base_data": state.get("base_data") or {},
    }
    return {
        "analysis_id": current.get("id"),
        "attachment": {**att, "form_data": form},
        "form_data": form,
        "kind": att.get("kind"),
        **extra,
    }


def save_plan_workbench(project_id: int, code: str, form_data: Dict[str, Any]) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find(state, code=code)
    if not att or att.get("kind") not in KIND_HANDLERS:
        raise ValueError("当前附件不在附件05–09工作台范围")
    kind = att.get("kind")
    merged = {**(att.get("form_data") or {}), **(form_data or {})}
    if kind == "labor_plan":
        merged = ensure_labor_plan(merged)
        # 若未显式传曲线，按阶段重算
        if "phases" in (form_data or {}) and "curve_points" not in (form_data or {}):
            merged["curve_points"] = _curve_from_phases(merged["phases"])
        # 峰值
        try:
            nums = [int(float(str(c.get("headcount") or 0))) for c in merged.get("curve_points") or []]
            merged["peak_headcount"] = max(nums) if nums else merged.get("peak_headcount") or ""
        except Exception:
            pass
    elif kind == "equipment_plan":
        merged = ensure_equipment_plan(merged)
    elif kind == "schedule":
        merged = ensure_schedule(merged)
    elif kind == "labor_agreement":
        merged = ensure_labor_agreement(merged, state.get("base_data") or {}, _org_master(project_id))
        # 回写 conditions.labor_checks
        lc = merged.get("legal_checklist") or {}
        patch_cond = {
            "labor_checks": {
                "use_employer_quota": lc.get("use_employer_quota"),
                "non_employer_work": lc.get("non_employer_work"),
                "accept_advance_deduct": lc.get("accept_advance_deduct"),
                "need_liability_boundary": lc.get("need_liability_boundary"),
            }
        }
        update_instructions(project_id, {"conditions": patch_cond})
    elif kind == "confidentiality":
        merged = ensure_confidentiality(merged, state.get("base_data") or {}, _org_master(project_id))

    status = _drafting(att.get("status"))
    if kind == "schedule" and not merged.get("confirmed"):
        status = "pending_p6" if att.get("status") == "pending_p6" else status
    update_instructions(project_id, {"attachment": {"code": code, "form_data": merged, "status": status}})
    return get_plan_workbench(project_id, code)


def generate_plan_draft(project_id: int, code: str) -> Dict[str, Any]:
    """按模板/进度依赖生成草稿数据。"""
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find(state, code=code)
    if not att or att.get("kind") not in KIND_HANDLERS:
        raise ValueError("当前附件不在附件05–09工作台范围")
    kind = att.get("kind")
    form = ensure_form_for_kind(project_id, att, state)
    base = state.get("base_data") or {}

    if kind == "labor_plan":
        # 给各阶段填入示意人数，便于曲线展示，用户再改
        demo = {"深化设计": 8, "材料进场": 15, "安装施工": 45, "调试验收": 12}
        for p in form["phases"]:
            if not str(p.get("headcount") or "").strip():
                ph = p.get("phase_name") or ""
                base_n = demo.get(ph, 10)
                # 按工种略分摊
                p["headcount"] = str(max(1, base_n // 3))
        form["curve_points"] = _curve_from_phases(form["phases"])
        form["peak_headcount"] = max(int(float(c["headcount"] or 0)) for c in form["curve_points"])
        form["notes"] = (form.get("notes") or "") or "已按常见机电分包阶段生成示意劳动力，请按实际工程量与进度调整峰值与工种组合。"
    elif kind == "equipment_plan":
        for r in form["equipment_rows"]:
            if not r.get("location"):
                r["location"] = "国内仓库/待发"
            if not r.get("source"):
                r["source"] = "由国内发货"
            if r.get("ready_now") is None:
                r["ready_now"] = True
        form["notes"] = (form.get("notes") or "") or "已按常见机电机具生成推荐清单，请核对数量、所在地、进场时间与调用可行性。"
    elif kind == "schedule":
        proj = base.get("project_name") or "本工程"
        for m in form["wbs_milestones"]:
            if not m.get("notes"):
                m["notes"] = f"{proj} — 待进度工程师在 P6/MSP 中细化"
        form["scope_covers"] = {"detail_design": True, "material_delivery": True, "construction_org": True}
        form["p6"]["export_status"] = "draft_ready"
        form["engineer_notes"] = form.get("engineer_notes") or "系统已生成 WBS/里程碑草稿，请导出后在 P6 或 Project 中编制正式计划并回传确认。"
    elif kind in ("labor_agreement", "confidentiality"):
        # 重新预填 + 刷新正文预览
        form["locked_body_preview"] = _extract_body(project_id, att)[:4000]
        master = _org_master(project_id)
        for f in form.get("fields") or []:
            if not str(f.get("value") or "").strip():
                f["value"] = _prefill_value(str(f.get("key") or ""), base, master)

    update_instructions(
        project_id,
        {"attachment": {"code": code, "form_data": form, "status": _drafting(att.get("status"))}},
    )
    return get_plan_workbench(project_id, code)


def render_labor_curve_svg(points: List[Dict[str, Any]]) -> str:
    if not points:
        return '<svg xmlns="http://www.w3.org/2000/svg" width="560" height="160"><text x="20" y="80" fill="#8c8c8c">暂无曲线数据</text></svg>'
    vals = []
    labels = []
    for p in points:
        try:
            vals.append(float(p.get("headcount") or 0))
        except Exception:
            vals.append(0.0)
        labels.append(str(p.get("period") or ""))
    w, h = 640, 220
    pad_l, pad_r, pad_t, pad_b = 48, 24, 24, 48
    max_v = max(vals) if vals else 1
    max_v = max_v or 1
    n = len(vals)
    usable_w = w - pad_l - pad_r
    usable_h = h - pad_t - pad_b
    coords = []
    for i, v in enumerate(vals):
        x = pad_l + (usable_w * i / max(n - 1, 1))
        y = pad_t + usable_h * (1 - v / max_v)
        coords.append((x, y))
    poly = " ".join(f"{x},{y}" for x, y in coords)
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">',
        '<rect width="100%" height="100%" fill="#fafafa"/>',
        f'<line x1="{pad_l}" y1="{pad_t}" x2="{pad_l}" y2="{h - pad_b}" stroke="#d9d9d9"/>',
        f'<line x1="{pad_l}" y1="{h - pad_b}" x2="{w - pad_r}" y2="{h - pad_b}" stroke="#d9d9d9"/>',
        f'<polyline fill="none" stroke="#1677ff" stroke-width="2.5" points="{poly}"/>',
    ]
    for (x, y), lab, v in zip(coords, labels, vals):
        parts.append(f'<circle cx="{x}" cy="{y}" r="4" fill="#1677ff"/>')
        parts.append(f'<text x="{x}" y="{h - 16}" text-anchor="middle" font-size="11" fill="#595959">{_esc(lab)}</text>')
        parts.append(f'<text x="{x}" y="{y - 8}" text-anchor="middle" font-size="11" fill="#262626">{int(v)}</text>')
    parts.append(f'<text x="8" y="{pad_t + 4}" font-size="11" fill="#8c8c8c">人</text>')
    parts.append("</svg>")
    return "\n".join(parts)


def _esc(s: str) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def validate_remaining_complete(project_id: int, att: Dict[str, Any], state: Dict[str, Any]) -> None:
    kind = att.get("kind")
    form = att.get("form_data") or {}
    if kind == "labor_plan":
        form = ensure_labor_plan(form)
        filled = [p for p in form["phases"] if str(p.get("headcount") or "").strip() not in ("", "0")]
        if len(filled) < 1:
            raise ValueError("请至少填写若干阶段的劳动力人数")
        cs = form.get("chinese_source") or {}
        if not form.get("voucher_attached") and not str(form.get("credential_notes") or "").strip():
            raise ValueError("请补充护照/签证/培训凭证说明，或勾选已附凭证")
        # soft: chinese plan
        _ = cs
    elif kind == "equipment_plan":
        form = ensure_equipment_plan(form)
        rows = [r for r in form["equipment_rows"] if str(r.get("name") or "").strip()]
        if not rows:
            raise ValueError("请至少填写一条设备/机具/周转材料")
        flags = form.get("confirm_flags") or {}
        missing = [k for k, lab in (
            ("source_explained", "来源说明"),
            ("mobilize_aligned_schedule", "进场时间与进度对齐"),
            ("voucher_prepared", "凭证已准备"),
            ("ready_confirmed", "可立即投入使用"),
        ) if not flags.get(k)]
        if missing:
            raise ValueError("请确认：" + "、".join(missing))
    elif kind == "schedule":
        form = ensure_schedule(form)
        ms = [m for m in form["wbs_milestones"] if str(m.get("name") or "").strip()]
        if not ms:
            raise ValueError("请生成或填写 WBS/里程碑草稿")
        scope = form.get("scope_covers") or {}
        if not all(scope.get(k) for k in ("detail_design", "material_delivery", "construction_org")):
            raise ValueError("进度计划须覆盖深化设计、材料进场、施工组织")
        if not form.get("confirmed"):
            raise ValueError("请确认进度工程师已在 P6/MSP 调整并回传/定稿")
    elif kind == "labor_agreement":
        cond = state.get("conditions") or {}
        if cond.get("labor_applicable") is False:
            return  # 将在 complete 里标 not_required
        if cond.get("labor_applicable") is None:
            raise ValueError("请先在条件判断中确认劳务协议是否适用")
        form = ensure_labor_agreement(form, state.get("base_data") or {}, _org_master(project_id))
        for f in form.get("fields") or []:
            if f.get("required") and not str(f.get("value") or "").strip():
                raise ValueError(f"请填写：{f.get('label') or f.get('key')}")
        lc = form.get("legal_checklist") or {}
        for k, lab in (
            ("use_employer_quota", "用工指标办证"),
            ("non_employer_work", "非甲方单位用工报备"),
            ("accept_advance_deduct", "代付扣款条款"),
            ("need_liability_boundary", "安全与赔偿责任边界"),
        ):
            if lc.get(k) is None:
                raise ValueError(f"请确认法务审查项：{lab}")
        if not lc.get("legal_reviewed"):
            raise ValueError("请勾选「法务已审查」")
    elif kind == "confidentiality":
        form = ensure_confidentiality(form, state.get("base_data") or {}, _org_master(project_id))
        for f in form.get("fields") or []:
            if f.get("required") and not str(f.get("value") or "").strip():
                raise ValueError(f"请填写：{f.get('label') or f.get('key')}")
        if not form.get("scope_confirm"):
            raise ValueError("请确认保密范围")
        if not form.get("duration_until_acceptance"):
            raise ValueError("请确认保密期限至竣工验收")
        if not form.get("supplier_nda_confirm"):
            raise ValueError("请特别确认不向供应商泄露")


def complete_remaining(project_id: int, code: str) -> Dict[str, Any]:
    from .attachment_workbench import get_attachment_workbench

    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find(state, code=code)
    if not att:
        raise ValueError("附件不存在")
    kind = att.get("kind")
    if kind not in KIND_HANDLERS:
        raise ValueError("非附件05–09")

    # 确保 form
    form = ensure_form_for_kind(project_id, att, state)
    att = {**att, "form_data": form}
    validate_remaining_complete(project_id, att, state)

    if kind == "labor_agreement" and (state.get("conditions") or {}).get("labor_applicable") is False:
        update_instructions(project_id, {"attachment": {"code": code, "status": "not_required", "form_data": form}})
        return get_attachment_workbench(project_id, code)

    status = "done"
    if kind == "labor_agreement":
        # 完成后也可保留 pending_legal 语义；这里直接 done
        status = "done"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form}})
    return get_attachment_workbench(project_id, code)


def export_remaining_docx(project_id: int, code: str) -> Tuple[bytes, str]:
    wb = get_plan_workbench(project_id, code)
    att = wb["attachment"]
    kind = att.get("kind")
    form = wb["form_data"]
    title = att.get("title") or code
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(15)

    base = wb.get("base_data") or {}
    if base.get("project_name"):
        doc.add_paragraph(f"项目：{base.get('project_name')}")

    if kind == "labor_plan":
        doc.add_paragraph("一、劳动力计划表（单位：人，按每班八小时）")
        _add_table(
            doc,
            ["序号", "阶段", "工种/级别", "人数", "备注"],
            [
                [str(i + 1), r.get("phase_name"), f"{r.get('trade')}/{r.get('level')}", r.get("headcount"), r.get("notes")]
                for i, r in enumerate(form.get("phases") or [])
            ],
        )
        doc.add_paragraph(f"峰值人数：{form.get('peak_headcount') or '—'}")
        doc.add_paragraph("二、中国工人安排计划")
        cs = form.get("chinese_source") or {}
        hq = cs.get("hq_dispatch") or {}
        op = cs.get("from_other_projects") or {}
        tr = cs.get("trained_certified") or {}
        doc.add_paragraph(f"· 总部选派：具备条件={hq.get('eligible')} 人数={hq.get('count')} {hq.get('note')}")
        doc.add_paragraph(
            f"· 其他项目调派：护照={op.get('passport_done')} 签证={op.get('visa_done')} 人数={op.get('count')} {op.get('note')}"
        )
        doc.add_paragraph(f"· 已培训考试：人数={tr.get('count')} {tr.get('note')}")
        doc.add_paragraph(f"凭证说明：{form.get('credential_notes') or '—'}")
        doc.add_paragraph("三、劳动力曲线（阶段汇总）")
        for c in form.get("curve_points") or []:
            mark = "（峰值）" if c.get("peak") else ""
            doc.add_paragraph(f"· {c.get('period')}: {c.get('headcount')} 人{mark}")
    elif kind == "equipment_plan":
        doc.add_paragraph("一、设备机具周转材料清单")
        _add_table(
            doc,
            ["序号", "名称", "规格型号", "年限", "数量", "所在地", "进场时间", "来源"],
            [
                [
                    str(i + 1),
                    r.get("name"),
                    r.get("spec"),
                    r.get("age_years"),
                    r.get("qty"),
                    r.get("location"),
                    r.get("mobilize_date"),
                    r.get("source"),
                ]
                for i, r in enumerate(form.get("equipment_rows") or [])
            ],
        )
        flags = form.get("confirm_flags") or {}
        doc.add_paragraph("二、确认项")
        for k, lab in (
            ("source_explained", "来源已说明"),
            ("mobilize_aligned_schedule", "进场与进度对齐"),
            ("voucher_prepared", "凭证已准备"),
            ("ready_confirmed", "可立即投入"),
        ):
            doc.add_paragraph(f"· {lab}: {'是' if flags.get(k) else '否'}")
        doc.add_paragraph(f"凭证备注：{form.get('voucher_note') or '—'}")
    elif kind == "schedule":
        doc.add_paragraph("（分包提供，包含深化设计计划、材料进场计划、施工组织计划）")
        doc.add_paragraph("一、WBS / 里程碑草稿")
        _add_table(
            doc,
            ["序号", "名称", "类型", "计划开始", "计划完成", "备注"],
            [
                [str(i + 1), m.get("name"), m.get("type"), m.get("planned_start"), m.get("planned_finish"), m.get("notes")]
                for i, m in enumerate(form.get("wbs_milestones") or [])
            ],
        )
        scope = form.get("scope_covers") or {}
        doc.add_paragraph(
            "覆盖："
            + f"深化设计={'是' if scope.get('detail_design') else '否'}；"
            + f"材料进场={'是' if scope.get('material_delivery') else '否'}；"
            + f"施工组织={'是' if scope.get('construction_org') else '否'}"
        )
        p6 = form.get("p6") or {}
        doc.add_paragraph(f"P6/MSP：导出={p6.get('export_status')} 导入={p6.get('import_status')} 文件={p6.get('file_name') or '—'}")
        doc.add_paragraph(f"工程师备注：{form.get('engineer_notes') or '—'}")
    elif kind == "labor_agreement":
        fmap = {f.get("key"): f.get("value") for f in form.get("fields") or []}
        body = form.get("locked_body_preview") or "（正文以原协议模板为准，系统不改写条款）"
        doc.add_paragraph("【说明】正文条款锁定，仅回填身份信息；以下为预览与填写值。")
        for f in form.get("fields") or []:
            doc.add_paragraph(f"{f.get('label')}: {f.get('value') or '________'}")
        doc.add_paragraph("")
        for line in body.splitlines():
            # 简易替换空白
            text = line
            if "乙方" in text and fmap.get("party_b"):
                text = re.sub(r"乙方：\s*_+", f"乙方：{fmap['party_b']}", text)
            doc.add_paragraph(text)
        lc = form.get("legal_checklist") or {}
        doc.add_paragraph("【法务确认】")
        doc.add_paragraph(f"审查人：{lc.get('legal_reviewer') or '—'} 已审查={'是' if lc.get('legal_reviewed') else '否'}")
        doc.add_paragraph(f"备注：{lc.get('legal_notes') or '—'}")
    elif kind == "confidentiality":
        fmap = {f.get("key"): str(f.get("value") or "").strip() for f in form.get("fields") or []}
        body = form.get("locked_body_preview") or ""
        filled = body
        if fmap.get("project_name"):
            filled = filled.replace("的 项目 工程", f"的 {fmap['project_name']} 工程")
            filled = filled.replace("有关 项目的信息", f"有关 {fmap['project_name']} 的信息")
        doc.add_paragraph("【说明】正文锁定，仅回填项目/投标人/日期。")
        for line in filled.splitlines():
            doc.add_paragraph(line)
        doc.add_paragraph("")
        doc.add_paragraph(f"投标人：{fmap.get('bidder_name') or '________'}")
        doc.add_paragraph(f"法定代表人：{fmap.get('legal_rep') or '________'}")
        doc.add_paragraph("签章：")
        doc.add_paragraph(f"日期：{fmap.get('sign_date') or '________'}")
        doc.add_paragraph(
            f"确认：范围={'是' if form.get('scope_confirm') else '否'}；"
            f"至竣工验收={'是' if form.get('duration_until_acceptance') else '否'}；"
            f"不对供应商泄密={'是' if form.get('supplier_nda_confirm') else '否'}"
        )

    if form.get("notes"):
        doc.add_paragraph(f"备注：{form.get('notes')}")
    for tip in form.get("sign_reminder") or []:
        doc.add_paragraph(f"· {tip}")

    buf = io.BytesIO()
    doc.save(buf)
    safe = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", title)
    return buf.getvalue(), f"{safe}_已填写.docx"


def _add_table(doc: Document, headers: List[str], rows: List[List[Any]]) -> None:
    table = doc.add_table(rows=1 + max(len(rows), 1), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    data = rows or [[""] * len(headers)]
    for r, row in enumerate(data, 1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = "" if val is None else str(val)
