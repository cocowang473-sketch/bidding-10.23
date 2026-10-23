"""附件03：根据附件04人员生成项目管理组织机构图。"""
from __future__ import annotations

import io
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import attachment_workbench as att_wb
from .instructions_analysis import (
    STATUS_LABEL,
    get_or_init_instructions,
    update_instructions,
)

# 默认汇报层级：项目经理为根，其余默认向项目经理汇报
ROOT_POSITION = "项目经理"
DEFAULT_CHILD_ORDER = ["设计人员", "商务人员", "现场技术人员"]


def _find_att(state: Dict[str, Any], *, code: Optional[str] = None, kind: Optional[str] = None):
    for a in state.get("attachments") or []:
        if code and a.get("code") == code:
            return a
        if kind and a.get("kind") == kind:
            return a
    return None


def _people_from_a04(project_id: int) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    state = get_or_init_instructions(project_id)["result"]
    p4 = _find_att(state, kind="personnel_plan")
    if not p4:
        return [], None
    form = att_wb.ensure_personnel_form_data(p4)
    people = [p for p in (form.get("people") or []) if str(p.get("name") or "").strip()]
    return people, p4


def _new_node(
    *,
    name: str,
    position: str,
    parent_id: Optional[str] = None,
    level: int = 0,
    person_id: str = "",
    deployable: str = "是",
) -> Dict[str, Any]:
    return {
        "id": str(uuid.uuid4()),
        "person_id": person_id or "",
        "name": name,
        "position": position,
        "parent_id": parent_id,
        "level": level,
        "deployable": deployable,
    }


def build_nodes_from_people(people: List[Dict[str, Any]], project_title: str = "") -> List[Dict[str, Any]]:
    """由附件04人员生成树节点：项目经理为根，其余向其汇报。"""
    if not people:
        return []

    root_candidates = [p for p in people if str(p.get("position") or "").strip() == ROOT_POSITION]
    others = [p for p in people if p not in root_candidates]

    nodes: List[Dict[str, Any]] = []
    title_id = None
    if project_title:
        title = _new_node(name=project_title, position="项目", parent_id=None, level=0)
        title_id = title["id"]
        nodes.append(title)

    root_parent = title_id
    root_level = 1 if title_id else 0
    root_ids: List[str] = []

    if root_candidates:
        for p in root_candidates:
            n = _new_node(
                name=str(p.get("name") or "").strip(),
                position=ROOT_POSITION,
                parent_id=root_parent,
                level=root_level,
                person_id=str(p.get("id") or ""),
                deployable=str(p.get("interview_ok") or "是"),
            )
            nodes.append(n)
            root_ids.append(n["id"])
    else:
        p0 = people[0]
        n = _new_node(
            name=str(p0.get("name") or "").strip(),
            position=str(p0.get("position") or "负责人"),
            parent_id=root_parent,
            level=root_level,
            person_id=str(p0.get("id") or ""),
        )
        nodes.append(n)
        root_ids.append(n["id"])
        others = people[1:]

    primary_root = root_ids[0]
    child_level = root_level + 1

    def sort_key(p: Dict[str, Any]):
        pos = str(p.get("position") or "").strip()
        try:
            return (DEFAULT_CHILD_ORDER.index(pos), pos, str(p.get("name") or ""))
        except ValueError:
            return (99, pos, str(p.get("name") or ""))

    for p in sorted(others, key=sort_key):
        nodes.append(
            _new_node(
                name=str(p.get("name") or "").strip(),
                position=str(p.get("position") or "").strip() or "管理人员",
                parent_id=primary_root,
                level=child_level,
                person_id=str(p.get("id") or ""),
                deployable=str(p.get("interview_ok") or "是"),
            )
        )
    return nodes


def _children_map(nodes: List[Dict[str, Any]]) -> Dict[Optional[str], List[Dict[str, Any]]]:
    m: Dict[Optional[str], List[Dict[str, Any]]] = {}
    for n in nodes:
        pid = n.get("parent_id")
        m.setdefault(pid, []).append(n)
    return m


def _esc(s: str) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def render_org_svg(nodes: List[Dict[str, Any]], *, width_hint: int = 960) -> str:
    """纯 SVG 组织架构图（层级树）。"""
    if not nodes:
        return (
            '<svg xmlns="http://www.w3.org/2000/svg" width="480" height="120">'
            '<text x="24" y="64" fill="#8c8c8c" font-size="14">暂无节点，请先在附件04填写人员后生成</text>'
            "</svg>"
        )

    by_id = {n["id"]: n for n in nodes}
    children = _children_map(nodes)
    roots = [n for n in nodes if not n.get("parent_id") or n.get("parent_id") not in by_id]
    if not roots:
        roots = nodes[:1]

    box_w, box_h = 150, 58
    h_gap, v_gap = 28, 48
    subtree_w: Dict[str, float] = {}

    def measure(n: Dict[str, Any]) -> float:
        kids = children.get(n["id"]) or []
        if not kids:
            subtree_w[n["id"]] = box_w
            return box_w
        w = sum(measure(k) for k in kids) + h_gap * (len(kids) - 1)
        subtree_w[n["id"]] = max(w, box_w)
        return subtree_w[n["id"]]

    for r in roots:
        measure(r)

    positions: Dict[str, Tuple[float, float]] = {}

    def place(n: Dict[str, Any], x_left: float, depth: int) -> None:
        w = subtree_w[n["id"]]
        cx = x_left + w / 2
        cy = 24 + depth * (box_h + v_gap)
        positions[n["id"]] = (cx, cy)
        kids = children.get(n["id"]) or []
        cursor = x_left + (w - (sum(subtree_w[k["id"]] for k in kids) + h_gap * max(len(kids) - 1, 0))) / 2
        for k in kids:
            kw = subtree_w[k["id"]]
            place(k, cursor, depth + 1)
            cursor += kw + h_gap

    total_w = sum(subtree_w[r["id"]] for r in roots) + h_gap * max(len(roots) - 1, 0)
    cursor = 40.0
    max_depth = 0
    for r in roots:
        place(r, cursor, 0)
        cursor += subtree_w[r["id"]] + h_gap
    for n in nodes:
        if n["id"] in positions:
            max_depth = max(max_depth, int(round((positions[n["id"]][1] - 24) / (box_h + v_gap))))

    svg_w = max(width_hint, int(total_w + 80))
    svg_h = int(24 + (max_depth + 1) * (box_h + v_gap) + 40)

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{svg_w}" height="{svg_h}" viewBox="0 0 {svg_w} {svg_h}">',
        "<defs>",
        "<style><![CDATA["
        ".box{fill:#fff;stroke:#1677ff;stroke-width:1.5}"
        ".pos{font:600 12px sans-serif;fill:#1677ff;text-anchor:middle}"
        ".name{font:14px sans-serif;fill:#262626;text-anchor:middle}"
        ".line{stroke:#91caff;stroke-width:1.5;fill:none}"
        "]]></style>",
        "</defs>",
        '<rect width="100%" height="100%" fill="#fafafa"/>',
    ]

    for n in nodes:
        pid = n.get("parent_id")
        if not pid or pid not in positions or n["id"] not in positions:
            continue
        x1, y1 = positions[pid]
        x2, y2 = positions[n["id"]]
        mid_y = (y1 + box_h / 2 + y2 - box_h / 2) / 2
        parts.append(
            f'<path class="line" d="M{x1},{y1 + box_h / 2} V{mid_y} H{x2} V{y2 - box_h / 2}"/>'
        )

    for n in nodes:
        if n["id"] not in positions:
            continue
        cx, cy = positions[n["id"]]
        x = cx - box_w / 2
        y = cy - box_h / 2
        pos = _esc(n.get("position") or "")
        name = _esc(n.get("name") or "")
        parts.append(
            f'<rect class="box" x="{x}" y="{y}" width="{box_w}" height="{box_h}" rx="6" ry="6"/>'
        )
        parts.append(f'<text class="pos" x="{cx}" y="{cy - 6}">{pos}</text>')
        parts.append(f'<text class="name" x="{cx}" y="{cy + 14}">{name}</text>')

    parts.append("</svg>")
    return "\n".join(parts)


def render_org_png(nodes: List[Dict[str, Any]]) -> bytes:
    """Pillow 绘制组织图 PNG，供 Word 嵌入。"""
    from PIL import Image, ImageDraw, ImageFont

    if not nodes:
        img = Image.new("RGB", (640, 160), "#fafafa")
        draw = ImageDraw.Draw(img)
        draw.text((24, 70), "暂无节点，请先在附件04填写人员后生成", fill="#8c8c8c")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()

    by_id = {n["id"]: n for n in nodes}
    children = _children_map(nodes)
    roots = [n for n in nodes if not n.get("parent_id") or n.get("parent_id") not in by_id]
    if not roots:
        roots = nodes[:1]

    box_w, box_h = 160, 64
    h_gap, v_gap = 32, 56
    subtree_w: Dict[str, float] = {}

    def measure(n: Dict[str, Any]) -> float:
        kids = children.get(n["id"]) or []
        if not kids:
            subtree_w[n["id"]] = box_w
            return box_w
        w = sum(measure(k) for k in kids) + h_gap * (len(kids) - 1)
        subtree_w[n["id"]] = max(w, box_w)
        return subtree_w[n["id"]]

    for r in roots:
        measure(r)

    positions: Dict[str, Tuple[float, float]] = {}

    def place(n: Dict[str, Any], x_left: float, depth: int) -> None:
        w = subtree_w[n["id"]]
        cx = x_left + w / 2
        cy = 40 + depth * (box_h + v_gap)
        positions[n["id"]] = (cx, cy)
        kids = children.get(n["id"]) or []
        cursor = x_left + (w - (sum(subtree_w[k["id"]] for k in kids) + h_gap * max(len(kids) - 1, 0))) / 2
        for k in kids:
            kw = subtree_w[k["id"]]
            place(k, cursor, depth + 1)
            cursor += kw + h_gap

    total_w = sum(subtree_w[r["id"]] for r in roots) + h_gap * max(len(roots) - 1, 0)
    cursor = 48.0
    max_depth = 0
    for r in roots:
        place(r, cursor, 0)
        cursor += subtree_w[r["id"]] + h_gap
    for n in nodes:
        if n["id"] in positions:
            max_depth = max(max_depth, int(round((positions[n["id"]][1] - 40) / (box_h + v_gap))))

    svg_w = max(720, int(total_w + 96))
    svg_h = int(40 + (max_depth + 1) * (box_h + v_gap) + 48)
    img = Image.new("RGB", (svg_w, svg_h), "#fafafa")
    draw = ImageDraw.Draw(img)
    try:
        font_pos = ImageFont.truetype("msyh.ttc", 13)
        font_name = ImageFont.truetype("msyh.ttc", 15)
    except Exception:
        font_pos = ImageFont.load_default()
        font_name = ImageFont.load_default()

    for n in nodes:
        pid = n.get("parent_id")
        if not pid or pid not in positions or n["id"] not in positions:
            continue
        x1, y1 = positions[pid]
        x2, y2 = positions[n["id"]]
        mid_y = (y1 + box_h / 2 + y2 - box_h / 2) / 2
        draw.line(
            [(x1, y1 + box_h / 2), (x1, mid_y), (x2, mid_y), (x2, y2 - box_h / 2)],
            fill="#91caff",
            width=2,
        )

    for n in nodes:
        if n["id"] not in positions:
            continue
        cx, cy = positions[n["id"]]
        x0, y0 = cx - box_w / 2, cy - box_h / 2
        draw.rounded_rectangle(
            [x0, y0, x0 + box_w, y0 + box_h],
            radius=8,
            outline="#1677ff",
            width=2,
            fill="#ffffff",
        )
        pos = str(n.get("position") or "")
        name = str(n.get("name") or "")
        draw.text((cx, cy - 12), pos, fill="#1677ff", font=font_pos, anchor="mm")
        draw.text((cx, cy + 12), name, fill="#262626", font=font_name, anchor="mm")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def default_confirm() -> Dict[str, Any]:
    return {
        "hierarchy_ok": None,
        "reporting_ok": None,
        "position_ok": None,
        "deployable_ok": None,
    }


def get_org_chart(project_id: int, code: str) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_att(state, code=code)
    if not att or att.get("kind") != "org_chart":
        raise ValueError("当前附件不是组织机构图")

    people, p4 = _people_from_a04(project_id)
    form = dict(att.get("form_data") or {})
    nodes = form.get("nodes") if isinstance(form.get("nodes"), list) else []
    base = state.get("base_data") or {}
    project_title = base.get("project_name") or ""
    svg = form.get("svg") or render_org_svg(nodes)
    return {
        "analysis_id": current.get("id"),
        "attachment": att,
        "nodes": nodes,
        "svg": svg,
        "confirm": form.get("confirm") or default_confirm(),
        "source": {
            "personnel_code": (p4 or {}).get("code"),
            "personnel_status": (p4 or {}).get("status"),
            "personnel_status_label": STATUS_LABEL.get(
                (p4 or {}).get("status") or "", (p4 or {}).get("status")
            ),
            "people_count": len(people),
            "project_title": project_title,
        },
        "has_chart": bool(nodes),
    }


def generate_org_chart(project_id: int, code: str) -> Dict[str, Any]:
    """根据附件04人员重新生成组织架构。"""
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_att(state, code=code)
    if not att or att.get("kind") != "org_chart":
        raise ValueError("当前附件不是组织机构图")

    people, p4 = _people_from_a04(project_id)
    if not people:
        raise ValueError("附件04尚无已填姓名的人员，请先完成管理人员配备表")

    base = state.get("base_data") or {}
    project_title = str(base.get("project_name") or "").strip()
    nodes = build_nodes_from_people(people, project_title=project_title)
    form = dict(att.get("form_data") or {})
    form["nodes"] = nodes
    form["svg"] = render_org_svg(nodes)
    form["generated_from"] = {
        "personnel_code": (p4 or {}).get("code"),
        "people_count": len(people),
    }
    form.setdefault("confirm", default_confirm())

    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form}})

    from . import config

    out_dir = Path(config.UPLOAD_DIR) / f"project_{project_id}" / "org_chart"
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "org_chart.png").write_bytes(render_org_png(nodes))
    (out_dir / "org_chart.svg").write_text(form["svg"], encoding="utf-8")

    return get_org_chart(project_id, code)


def save_org_chart(
    project_id: int,
    code: str,
    *,
    nodes: Optional[List[Dict[str, Any]]] = None,
    confirm: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_att(state, code=code)
    if not att or att.get("kind") != "org_chart":
        raise ValueError("当前附件不是组织机构图")

    form = dict(att.get("form_data") or {})
    if nodes is not None:
        cleaned = []
        for n in nodes:
            if not isinstance(n, dict):
                continue
            cleaned.append(
                {
                    "id": str(n.get("id") or uuid.uuid4()),
                    "person_id": str(n.get("person_id") or ""),
                    "name": str(n.get("name") or "").strip(),
                    "position": str(n.get("position") or "").strip(),
                    "parent_id": n.get("parent_id") or None,
                    "level": int(n.get("level") or 0),
                    "deployable": str(n.get("deployable") or ""),
                }
            )
        by_id = {n["id"]: n for n in cleaned}

        def depth(nid: Optional[str], guard: int = 0) -> int:
            if not nid or nid not in by_id or guard > 20:
                return 0
            pid = by_id[nid].get("parent_id")
            if not pid or pid not in by_id:
                return 0
            return 1 + depth(pid, guard + 1)

        for n in cleaned:
            n["level"] = depth(n["id"])
        form["nodes"] = cleaned
        form["svg"] = render_org_svg(cleaned)
    if confirm is not None:
        form["confirm"] = {**default_confirm(), **(form.get("confirm") or {}), **confirm}

    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form}})
    return get_org_chart(project_id, code)


def validate_org_chart_complete(att: Dict[str, Any]) -> None:
    form = att.get("form_data") or {}
    nodes = form.get("nodes") or []
    named = [n for n in nodes if str(n.get("name") or "").strip()]
    if len(named) < 1:
        raise ValueError("请先根据附件04生成组织机构图")
    confirm = form.get("confirm") or {}
    labels = {
        "hierarchy_ok": "层级关系",
        "reporting_ok": "汇报关系",
        "position_ok": "岗位设置",
        "deployable_ok": "人员真实可投入",
    }
    missing = [lab for k, lab in labels.items() if confirm.get(k) is not True]
    if missing:
        raise ValueError("请确认：" + "、".join(missing))


def export_org_chart_docx(project_id: int, code: str) -> Tuple[bytes, str]:
    data = get_org_chart(project_id, code)
    att = data["attachment"]
    nodes = data.get("nodes") or []
    if not nodes:
        raise ValueError("请先生成组织机构图")

    doc = Document()
    title = att.get("title") or "分包商项目管理组织机构图"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    src = data.get("source") or {}
    if src.get("project_title"):
        doc.add_paragraph(f"项目：{src['project_title']}")
    doc.add_paragraph(f"人员来源：附件04（{src.get('people_count', 0)} 人）")

    png = render_org_png(nodes)
    doc.add_picture(io.BytesIO(png), width=Inches(6.2))

    doc.add_paragraph("")
    doc.add_paragraph("【节点明细】")
    for n in sorted(nodes, key=lambda x: (int(x.get("level") or 0), str(x.get("position") or ""))):
        indent = "　" * int(n.get("level") or 0)
        parent = ""
        if n.get("parent_id"):
            parent_name = next((x.get("name") for x in nodes if x.get("id") == n.get("parent_id")), "")
            if parent_name:
                parent = f"（汇报给：{parent_name}）"
        doc.add_paragraph(
            f"{indent}· [{n.get('position') or ''}] {n.get('name') or ''} {parent} 可投入:{n.get('deployable') or '—'}"
        )

    confirm = data.get("confirm") or {}
    doc.add_paragraph("")
    doc.add_paragraph("【确认项】")
    for k, lab in (
        ("hierarchy_ok", "层级"),
        ("reporting_ok", "汇报关系"),
        ("position_ok", "岗位设置"),
        ("deployable_ok", "真实可投入"),
    ):
        v = confirm.get(k)
        mark = "已确认" if v is True else ("否" if v is False else "未确认")
        doc.add_paragraph(f"· {lab}：{mark}")

    buf = io.BytesIO()
    doc.save(buf)
    safe = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", title)
    return buf.getvalue(), f"{safe}_组织机构图.docx"
