"""投标须知附件工作台：逐个附件处理。

附件01（授权委托书）：
1. 条件判断是否已授权备案；
2. 分析模板需填字段；
3. 企业/法人/代理人主数据预填；
4. 表格编辑后导出 Word 供盖章。
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from . import db
from .ingest import absolute_doc_path, list_documents
from .instructions_analysis import (
    STATUS_LABEL,
    get_or_init_instructions,
    latest_instructions,
    update_instructions,
)
from .llm import call_dashscope_json
from .text_extract import extract_text

FIELD_ANALYZE_SYSTEM = """你是招标投标附件模板解析专家。
任务：从投标附件模板原文中，识别所有需要人工填写的空白字段（括号空白、下划线、冒号后空位等）。
硬性规则：
1. 只识别模板中明确的填空处，不要编造无关字段；模板正文固定条款不要当成字段。
2. 每个字段给出稳定的 key（英文蛇形）、中文 label、source（enterprise|legal_rep|agent|project|quote|manual）。
3. required 标明是否必填。
4. evidence 必须原样摘录模板中的空白片段（含空格、下划线、括号标签），例如「        (姓名)」或「____（单位名称）」，不要改写成「我（签字人姓名）」这类归纳句。
5. 日期类空白（如「20 年 月 日」「日期： 年 月 日」）合并为【一个】字段，key 用 auth_valid_until 或 sign_date，hint 写 YYYY-MM-DD；不要拆成年/月/日三个字段。
6. 「姓名/身份证号」合并为一个字段 agent_name_id，hint 写 姓名/身份证号。
7. 只输出 JSON。
返回：
{
  "fields": [
    {
      "key": "legal_rep_name",
      "label": "姓名",
      "source": "legal_rep",
      "required": true,
      "hint": "填写说明",
      "evidence": "        (姓名)"
    }
  ],
  "sign_reminder": ["代表人签字", "公司盖章"]
}
"""

# 模板抽不出来时的兜底字段（对齐《附件01 授权委托书》常见空白）
POA_FALLBACK_FIELDS: List[Dict[str, Any]] = [
    {"key": "legal_rep_name", "label": "法定代表人姓名", "source": "legal_rep", "required": True, "hint": "我 (姓名)系…的法定代表人", "evidence": "我 (姓名)系"},
    {"key": "bidder_name", "label": "投标人名称", "source": "enterprise", "required": True, "hint": "投标人全称", "evidence": "(投标人名称)"},
    {"key": "agent_unit_name", "label": "被授权单位名称", "source": "agent", "required": True, "hint": "代理人所属单位", "evidence": "（单位名称）"},
    {"key": "agent_name_id", "label": "代理人姓名/身份证号", "source": "agent", "required": True, "hint": "格式：姓名/身份证号", "evidence": "(姓名/身份证号)"},
    {"key": "auth_valid_until", "label": "授权有效期至", "source": "manual", "required": True, "hint": "格式：YYYY-MM-DD 或 YYYY/MM/DD", "evidence": "有效期至20 年 月 日"},
    {"key": "agent_gender", "label": "代理人性别", "source": "agent", "required": False, "hint": "男/女", "evidence": "性别："},
    {"key": "agent_age", "label": "代理人年龄", "source": "agent", "required": False, "hint": "周岁", "evidence": "年龄："},
    {"key": "agent_company", "label": "代理人所在单位", "source": "agent", "required": False, "hint": "单位", "evidence": "单位："},
    {"key": "agent_department", "label": "代理人所在部门", "source": "agent", "required": False, "hint": "所在部门", "evidence": "部门："},
    {"key": "agent_position", "label": "代理人职务", "source": "agent", "required": False, "hint": "职务", "evidence": "职务："},
    {"key": "sign_date", "label": "签署日期", "source": "manual", "required": True, "hint": "格式：YYYY-MM-DD 或 YYYY/MM/DD", "evidence": "日期： 年 月 日"},
]

POA_SIGN_REMINDERS = [
    "确认代理权限范围（投标、开标、评标、合同谈判）",
    "确认授权有效期",
    "代理人无转委托权",
    "法定代表人签字并加盖法人章、公司章后方可生效",
]

# 附件02 投标承诺函兜底字段
BID_FALLBACK_FIELDS: List[Dict[str, Any]] = [
    {"key": "signer_name", "label": "签字人姓名", "source": "legal_rep", "required": True, "hint": "我（签字人姓名）", "evidence": "我（签字人姓名）"},
    {"key": "company_name", "label": "公司名称", "source": "enterprise", "required": True, "hint": "以（公司名称）名义", "evidence": "（公司名称）"},
    {"key": "employer_name", "label": "招标人/递交对象", "source": "project", "required": True, "hint": "向…递交了", "evidence": "中建阿尔及利亚股份有限公司"},
    {"key": "project_name", "label": "项目/工程名称", "source": "project", "required": True, "hint": "…的投标书", "evidence": "递交了…的投标书"},
    {"key": "amount_lower", "label": "工程总价（小写）", "source": "quote", "required": True, "hint": "数字金额，如 1234567.89", "evidence": "工程总价：（小写）"},
    {"key": "amount_upper", "label": "工程总价（大写）", "source": "quote", "required": True, "hint": "可点按钮由小写自动生成后核对", "evidence": "（大写） 第纳尔"},
    {"key": "currency", "label": "币种", "source": "project", "required": True, "hint": "如 DA / 阿尔及利亚第纳尔", "evidence": "DA"},
    {"key": "sign_date", "label": "签署日期", "source": "manual", "required": True, "hint": "YYYY-MM-DD", "evidence": "日 期："},
]

BID_SIGN_REMINDERS = [
    "确认承诺内容与现场考察事实一致（已考察 / 资料评估 / 需澄清）",
    "校验小写金额与大写金额一致",
    "代表人签字、公司盖章后方可作为投标文件组成部分",
]

# 附件04 管理人员配备计划表（对齐模板列）
PERSONNEL_DEFAULT_POSITIONS = ["项目经理", "设计人员", "商务人员", "现场技术人员"]
PERSONNEL_COLUMNS = [
    {"key": "seq", "label": "序号", "width": 70},
    {"key": "position", "label": "岗位", "width": 120},
    {"key": "name", "label": "人员姓名", "width": 110},
    {"key": "birth_date", "label": "出生日期", "width": 120},
    {"key": "education_major", "label": "学历和专业", "width": 160},
    {"key": "work_years", "label": "工作年限", "width": 90},
    {"key": "hospital_exp", "label": "医院经验", "width": 90},
    {"key": "algeria_exp", "label": "阿尔及利亚经验", "width": 120},
    {"key": "local_subcontract_exp", "label": "属地分包经验", "width": 110},
    {"key": "overseas_procurement_exp", "label": "境外采购经验", "width": 110},
    {"key": "interview_ok", "label": "可参加面试", "width": 100},
]

PERSONNEL_PARSE_SYSTEM = """你是投标人员资料解析专家。
任务：从用户上传的个人简介、简历或人员配备表格中，识别管理人员信息行。
硬性规则：
1. 只输出文档中能支撑的信息，不要编造；不确定的字段留空字符串。
2. 若是表格，按行提取；若是个人简介，可只提取 1 人，岗位可从文中推断或留空。
3. position 尽量映射为：项目经理 / 设计人员 / 商务人员 / 现场技术人员；无法映射则保留原文岗位。
4. work_years 用数字或「X年」字符串均可。
5. 经验类字段用「有/无/不详」或简短说明。
6. 只输出 JSON。
返回：
{
  "people": [
    {
      "position": "项目经理",
      "name": "张三",
      "birth_date": "1985-03-01",
      "education_major": "本科 / 机电工程",
      "work_years": "15",
      "hospital_exp": "有",
      "algeria_exp": "不详",
      "local_subcontract_exp": "",
      "overseas_procurement_exp": "",
      "interview_ok": "是",
      "evidence": "原文短摘录"
    }
  ]
}
"""


def _new_person_row(position: str = "", seq: int = 1, **extra: Any) -> Dict[str, Any]:
    row = {
        "id": str(uuid.uuid4()),
        "seq": seq,
        "position": position or "",
        "name": "",
        "birth_date": "",
        "education_major": "",
        "work_years": "",
        "hospital_exp": "",
        "algeria_exp": "",
        "local_subcontract_exp": "",
        "overseas_procurement_exp": "",
        "interview_ok": "",
    }
    for k, v in extra.items():
        if k in row and v is not None:
            row[k] = str(v).strip() if not isinstance(v, (int, float)) else v
        elif k not in ("id",) and v is not None:
            row[k] = v
    return row


def default_personnel_people() -> List[Dict[str, Any]]:
    return [_new_person_row(pos, i) for i, pos in enumerate(PERSONNEL_DEFAULT_POSITIONS, 1)]


def _normalize_people(rows: List[Any]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for i, raw in enumerate(rows or [], 1):
        if not isinstance(raw, dict):
            continue
        row = _new_person_row(str(raw.get("position") or ""), int(raw.get("seq") or i))
        if raw.get("id"):
            row["id"] = str(raw["id"])
        for k in (
            "name",
            "birth_date",
            "education_major",
            "work_years",
            "hospital_exp",
            "algeria_exp",
            "local_subcontract_exp",
            "overseas_procurement_exp",
            "interview_ok",
        ):
            if raw.get(k) is not None:
                row[k] = str(raw.get(k) or "").strip()
        if raw.get("evidence"):
            row["evidence"] = str(raw.get("evidence") or "").strip()
        row["seq"] = i
        out.append(row)
    return out


def ensure_personnel_form_data(att: Dict[str, Any]) -> Dict[str, Any]:
    """保证 personnel_plan 附件有 people 表结构。"""
    form = dict(att.get("form_data") or {})
    people = form.get("people")
    if not isinstance(people, list) or not people:
        form["people"] = default_personnel_people()
    else:
        form["people"] = _normalize_people(people)
    form["columns"] = PERSONNEL_COLUMNS
    form.setdefault("import_meta", {})
    att["form_data"] = form
    return form


def get_personnel_table(project_id: int, code: str) -> Dict[str, Any]:
    wb = get_attachment_workbench(project_id, code)
    att = wb["attachment"]
    if att.get("kind") != "personnel_plan":
        raise ValueError("当前附件不是管理人员配备计划")
    form = ensure_personnel_form_data(att)
    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(
        project_id,
        {"attachment": {"code": code, "form_data": form, "status": status}},
    )
    # 重新读取，保证与库一致
    wb = get_attachment_workbench(project_id, code)
    att = wb["attachment"]
    form = ensure_personnel_form_data(att)
    return {
        **wb,
        "attachment": {**att, "form_data": form},
        "people": form["people"],
        "columns": PERSONNEL_COLUMNS,
    }


def save_personnel_table(
    project_id: int,
    code: str,
    people: List[Dict[str, Any]],
    *,
    notes: Optional[str] = None,
) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError(f"未找到附件 {code}")
    if att.get("kind") != "personnel_plan":
        raise ValueError("当前附件不是管理人员配备计划")
    form = ensure_personnel_form_data(att)
    form["people"] = _normalize_people(people)
    if notes is not None:
        form["notes"] = notes
    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form}})
    return get_personnel_table(project_id, code)


def _extract_import_text(path: Path, max_chars: int = 40000) -> Tuple[str, str]:
    """支持简介/表格：docx/pdf/txt/doc/xlsx/csv。"""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), data_only=True, read_only=True)
            lines = []
            for ws in wb.worksheets[:3]:
                lines.append(f"# sheet:{ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() if c is not None else "" for c in row]
                    if any(cells):
                        lines.append("\t".join(cells))
            wb.close()
            text = "\n".join(lines)[:max_chars]
            return text, "xlsx"
        except Exception as e:
            return "", f"error:xlsx:{type(e).__name__}"
    text, method = extract_text(path, max_chars=max_chars)
    return text, method


def import_personnel_file(
    project_id: int,
    code: str,
    file_bytes: bytes,
    filename: str,
    *,
    mode: str = "merge",
) -> Dict[str, Any]:
    """上传个人简介或人员表，自动识别后写入 people。mode=merge|replace|append。"""
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError(f"未找到附件 {code}")
    if att.get("kind") != "personnel_plan":
        raise ValueError("当前附件不是管理人员配备计划")

    from . import config

    root = Path(config.UPLOAD_DIR) / f"project_{project_id}" / "personnel_imports"
    root.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", filename or "import.bin")
    path = root / f"{uuid.uuid4().hex[:8]}_{safe_name}"
    path.write_bytes(file_bytes)

    text, method = _extract_import_text(path)
    if len(text.strip()) < 20:
        raise ValueError(f"无法读取导入文件内容（method={method}），请换 PDF/Word/Excel/TXT")

    parsed: List[Dict[str, Any]] = []
    llm_error = None
    try:
        prompt = (
            f"文件名: {filename}\n"
            f"提取方法: {method}\n"
            f"请识别下列内容中的管理人员信息（表格或个人简介均可）：\n\n{text[:16000]}"
        )
        result = call_dashscope_json(prompt, PERSONNEL_PARSE_SYSTEM, temperature=0.1)
        for p in result.get("people") or []:
            if isinstance(p, dict) and (p.get("name") or p.get("position")):
                parsed.append(p)
    except Exception as e:
        llm_error = f"{type(e).__name__}: {str(e)[:200]}"

    if not parsed and not llm_error:
        llm_error = "模型未识别到有效人员行"

    form = ensure_personnel_form_data(att)
    existing = form.get("people") or default_personnel_people()
    incoming = _normalize_people(parsed) if parsed else []

    mode = (mode or "merge").lower()
    if mode == "replace" and incoming:
        people = incoming
    elif mode == "append" and incoming:
        people = _normalize_people(list(existing) + incoming)
    else:
        # merge：按岗位优先回填空行，其余追加
        people = [dict(r) for r in existing]
        used = set()
        for inc in incoming:
            matched = False
            pos = (inc.get("position") or "").strip()
            for i, row in enumerate(people):
                if i in used:
                    continue
                if pos and row.get("position") == pos and not (row.get("name") or "").strip():
                    for k, v in inc.items():
                        if k in ("id", "seq"):
                            continue
                        if v:
                            row[k] = v
                    used.add(i)
                    matched = True
                    break
            if not matched and (inc.get("name") or inc.get("position")):
                people.append(inc)
        people = _normalize_people(people)

    form["people"] = people
    form["import_meta"] = {
        "filename": filename,
        "method": method,
        "mode": mode,
        "parsed_count": len(incoming),
        "llm_error": llm_error,
        "saved_as": str(path.name),
    }
    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form}})
    out = get_personnel_table(project_id, code)
    out["import_meta"] = form["import_meta"]
    out["parsed_preview"] = incoming
    return out


def _export_personnel_docx(att: Dict[str, Any], people: List[Dict[str, Any]]) -> Tuple[bytes, str]:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    title = att.get("title") or "施工管理人员配备计划及主要管理人员履历"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(14)

    headers = ["序号", "岗位", "人员姓名", "出生日期", "学历和专业", "工作年限"]
    keys = ["seq", "position", "name", "birth_date", "education_major", "work_years"]
    rows = people or [_new_person_row()]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r, person in enumerate(rows, 1):
        for c, k in enumerate(keys):
            val = person.get(k, "")
            table.rows[r].cells[c].text = str(val if val is not None else "")

    doc.add_paragraph("")
    doc.add_paragraph("【经验与面试核对】")
    for person in rows:
        name = person.get("name") or "（未填姓名）"
        pos = person.get("position") or ""
        bits = [
            f"医院经验:{person.get('hospital_exp') or '—'}",
            f"阿尔及利亚:{person.get('algeria_exp') or '—'}",
            f"属地分包:{person.get('local_subcontract_exp') or '—'}",
            f"境外采购:{person.get('overseas_procurement_exp') or '—'}",
            f"可面试:{person.get('interview_ok') or '—'}",
        ]
        doc.add_paragraph(f"· {pos} {name} — " + "；".join(bits))

    buf = io.BytesIO()
    doc.save(buf)
    safe = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", title)
    return buf.getvalue(), f"{safe}_配备计划.docx"


def _master_path(project_id: int) -> Path:
    from . import config

    root = Path(config.UPLOAD_DIR) / f"project_{project_id}"
    root.mkdir(parents=True, exist_ok=True)
    return root / "org_master.json"


def get_org_master(project_id: int) -> Dict[str, Any]:
    path = _master_path(project_id)
    if path.exists():
        return db.loads(path.read_text(encoding="utf-8"), {})
    # 默认空主数据 + 从须知/公告带一点项目信息
    from .instructions_analysis import latest_instructions

    base = {}
    inst = latest_instructions(project_id)
    if inst and inst.get("result"):
        base = (inst["result"].get("base_data") or {}).copy()
    return {
        "enterprise": {
            "bidder_name": base.get("bidder_name") or "",
            "credit_code": "",
            "address": "",
            "phone": "",
        },
        "legal_rep": {
            "legal_rep_name": base.get("legal_representative") or "",
            "legal_rep_id_no": "",
            "legal_rep_title": "法定代表人",
        },
        "agent": {
            "agent_org_name": base.get("bidder_name") or "",
            "agent_name": base.get("agent_name") or "",
            "agent_id_no": "",
            "agent_gender": "",
            "agent_age": "",
            "agent_dept": "",
            "agent_title": "",
        },
        "project": {
            "project_name": base.get("project_name") or "",
            "employer_name": base.get("employer") or "",
            "scope": base.get("scope") or "",
        },
    }


def save_org_master(project_id: int, data: Dict[str, Any]) -> Dict[str, Any]:
    cur = get_org_master(project_id)
    for bucket in ("enterprise", "legal_rep", "agent", "project"):
        if bucket in data and isinstance(data[bucket], dict):
            cur[bucket] = {**(cur.get(bucket) or {}), **data[bucket]}
    path = _master_path(project_id)
    path.write_text(db.dumps(cur), encoding="utf-8")
    return cur


def _find_attachment(state: Dict[str, Any], code: Optional[str] = None, kind: Optional[str] = None) -> Optional[Dict[str, Any]]:
    for a in state.get("attachments") or []:
        if code and a.get("code") == code:
            return a
        if kind and a.get("kind") == kind:
            return a
    return None


def current_work_attachment(state: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """返回第一个「纳入且未完成」的附件（仅作进度提示，不限制操作顺序）。"""
    for a in state.get("attachments") or []:
        if a.get("included") is False:
            continue
        if a.get("status") in ("done", "not_required", "skipped"):
            continue
        return a
    return None


def _clean_template_text(raw: str) -> str:
    # 去掉 OLE 噪声，保留中文/常见标点与数字字母
    chars = []
    for ch in raw or "":
        o = ord(ch)
        if 0x4E00 <= o <= 0x9FFF or ch in "\n\r\t ，。；：、（）()《》[]【】—-/.%0123456789":
            chars.append(ch)
        elif ch.isascii() and (ch.isalnum() or ch in " \n\r\t.,;:()/-_%"):
            chars.append(ch)
    text = "".join(chars)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_attachment_text(project_id: int, att: Dict[str, Any]) -> Tuple[str, str]:
    path = absolute_doc_path(project_id, att["rel_path"])
    text, method = extract_text(path, max_chars=20000)
    text = _clean_template_text(text)
    return text, method


def analyze_attachment_fields(project_id: int, code: str) -> Dict[str, Any]:
    """分析模板需填字段，并写入该附件 form_data.fields。"""
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError(f"未找到附件 {code}")

    kind = att.get("kind") or ""
    fallback = BID_FALLBACK_FIELDS if kind == "bid_commitment" else POA_FALLBACK_FIELDS
    default_reminders = BID_SIGN_REMINDERS if kind == "bid_commitment" else POA_SIGN_REMINDERS

    text, method = _extract_attachment_text(project_id, att)
    fields: List[Dict[str, Any]] = []
    sign_reminder = list(default_reminders)
    llm_error = None

    if len(text) < 80:
        fields = [dict(f) for f in fallback]
        llm_error = f"模板文本过短/不可读（method={method}），已使用标准字段兜底"
    else:
        try:
            prompt = (
                f"附件文件名: {att.get('filename')}\n"
                f"附件类型: {kind} / {att.get('title')}\n"
                f"请识别下列模板中所有需填写字段（只认模板空白，不要改写固定条款）：\n\n{text[:12000]}"
            )
            result = call_dashscope_json(prompt, FIELD_ANALYZE_SYSTEM, temperature=0.1)
            for f in result.get("fields") or []:
                if not isinstance(f, dict):
                    continue
                key = str(f.get("key") or "").strip()
                label = str(f.get("label") or "").strip()
                if not key or not label:
                    continue
                fields.append(
                    {
                        "key": key,
                        "label": label,
                        "source": f.get("source") or "manual",
                        "required": bool(f.get("required", True)),
                        "hint": f.get("hint") or "",
                        "evidence": f.get("evidence") or "",
                        "value": "",
                    }
                )
            if result.get("sign_reminder"):
                sign_reminder = [str(x) for x in result["sign_reminder"] if x]
            if not fields:
                fields = [dict(f) for f in fallback]
                llm_error = "模型未返回有效 fields，已使用兜底字段"
        except Exception as e:
            fields = [dict(f) for f in fallback]
            llm_error = f"{type(e).__name__}: {str(e)[:160]}"

    # 保留用户已填 value
    old_vals = {}
    for f in ((att.get("form_data") or {}).get("fields") or []):
        if f.get("key") is not None:
            old_vals[f["key"]] = f.get("value")
    for f in fields:
        if f["key"] in old_vals and old_vals[f["key"]] not in (None, ""):
            f["value"] = old_vals[f["key"]]

    # 主数据 + 报价预填空值
    master = get_org_master(project_id)
    fields = prefill_fields_from_master(fields, master, base_data=state.get("base_data") or {})

    form_data = {
        **(att.get("form_data") or {}),
        "fields": fields,
        "sign_reminder": sign_reminder,
        "template_method": method,
        "template_text_len": len(text),
        "template_preview": text[:1500],
        "template_body": text,  # 导出时基于原模板正文回填，不另起生成
        "llm_error": llm_error,
        "analyzed_at": date.today().isoformat(),
    }
    status = att.get("status")
    if status in ("pending_decision", "to_prepare"):
        status = "drafting"
    update_instructions(
        project_id,
        {"attachment": {"code": code, "status": status, "form_data": form_data}},
    )
    # 若是未备案进入编制，状态改为 drafting
    if kind == "power_of_attorney":
        cond = (latest_instructions(project_id) or {}).get("result", {}).get("conditions") or {}
        if cond.get("auth_filing") == "not_filed":
            update_instructions(project_id, {"attachment": {"code": code, "status": "drafting", "form_data": form_data}})
        elif cond.get("auth_filing") == "filed":
            update_instructions(
                project_id,
                {"attachment": {"code": code, "status": "not_required", "form_data": form_data}},
            )

    return get_attachment_workbench(project_id, code)


def prefill_fields_from_master(
    fields: List[Dict[str, Any]],
    master: Dict[str, Any],
    base_data: Optional[Dict[str, Any]] = None,
) -> List[Dict[str, Any]]:
    flat: Dict[str, Any] = {}
    for bucket in ("enterprise", "legal_rep", "agent", "project"):
        flat.update(master.get(bucket) or {})
    base_data = base_data or {}
    # 别名映射
    aliases = {
        "bidder_name": flat.get("bidder_name") or base_data.get("bidder_name"),
        "company_name": flat.get("bidder_name") or base_data.get("bidder_name"),
        "legal_rep_name": flat.get("legal_rep_name") or base_data.get("legal_representative"),
        "signer_name": flat.get("legal_rep_name")
        or flat.get("agent_name")
        or base_data.get("legal_representative")
        or base_data.get("agent_name"),
        "agent_name": flat.get("agent_name"),
        "agent_id_no": flat.get("agent_id_no"),
        "agent_id_card": flat.get("agent_id_no"),
        "agent_name_id": (
            f"{flat.get('agent_name')}/{flat.get('agent_id_no')}"
            if flat.get("agent_name") and flat.get("agent_id_no")
            else (flat.get("agent_name") or "")
        ),
        "agent_unit_name": flat.get("agent_org_name") or flat.get("bidder_name"),
        "agent_org_name": flat.get("agent_org_name") or flat.get("bidder_name"),
        "agent_company": flat.get("agent_org_name") or flat.get("bidder_name"),
        "agent_gender": flat.get("agent_gender"),
        "agent_age": flat.get("agent_age"),
        "agent_dept": flat.get("agent_dept"),
        "agent_department": flat.get("agent_dept"),
        "agent_title": flat.get("agent_title"),
        "agent_position": flat.get("agent_title"),
        "project_name": flat.get("project_name") or base_data.get("project_name"),
        "employer_name": flat.get("employer_name") or flat.get("employer") or base_data.get("employer"),
        "employer": flat.get("employer_name") or flat.get("employer") or base_data.get("employer"),
        "currency": base_data.get("currency") or "DA",
        "amount_lower": _normalize_amount_str(base_data.get("locked_quote")),
        "total_price_lower": _normalize_amount_str(base_data.get("locked_quote")),
        "quote_amount_digits": _normalize_amount_str(base_data.get("locked_quote")),
        "amount_upper": "",
        "total_price_upper": "",
        "quote_amount_words": "",
        "sign_date": date.today().isoformat(),
        "date": date.today().isoformat(),
        "auth_valid_until": "",
    }
    # 若有小写金额，自动生成大写
    if aliases.get("amount_lower"):
        try:
            up = amount_to_cn_upper(aliases["amount_lower"])
            aliases["amount_upper"] = up
            aliases["total_price_upper"] = up
            aliases["quote_amount_words"] = up
        except Exception:
            pass

    for f in fields:
        if f.get("value") not in (None, ""):
            continue
        key = f.get("key")
        if key in aliases and aliases[key]:
            f["value"] = aliases[key]
    return fields


def _normalize_amount_str(raw: Any) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s:
        return ""
    # 抽出数字
    m = re.search(r"(\d+(?:[.,]\d+)?)", s.replace(",", "").replace(" ", ""))
    return m.group(1) if m else ""


def amount_to_cn_upper(amount: Any) -> str:
    """阿拉伯数字金额 → 中文大写（第纳尔/人民币通用财务大写）。"""
    s = _normalize_amount_str(amount)
    if not s:
        return ""
    try:
        num = float(s)
    except Exception as e:
        raise ValueError(f"无法解析金额: {amount}") from e
    if num < 0:
        raise ValueError("金额不能为负")

    digits = "零壹贰叁肆伍陆柒捌玖"
    int_units = ["", "拾", "佰", "仟"]
    dec_units = ["角", "分"]
    big_units = ["", "万", "亿"]

    # 保留两位小数
    whole = int(round(num * 100))
    yuan, fen = divmod(whole, 100)
    jiao, fen = divmod(fen, 10)

    def _sec(n: int) -> str:
        """千以内段：禁止段首多余「零」（如 0250 → 贰佰伍拾，而非 零贰佰伍拾）。"""
        if n == 0:
            return ""
        s4 = f"{n:04d}"
        out: List[str] = []
        pending_zero = False
        started = False
        for i, ch in enumerate(s4):
            d = int(ch)
            u = int_units[3 - i]
            if d == 0:
                if started:
                    pending_zero = True
                continue
            if pending_zero and started:
                out.append("零")
            pending_zero = False
            started = True
            out.append(digits[d] + u)
        return "".join(out)

    if yuan == 0:
        int_part = "零"
    else:
        parts: List[str] = []
        unit_i = 0
        n = yuan
        pending_zero = False
        while n > 0 and unit_i < len(big_units):
            n, sec = divmod(n, 10000)
            if sec:
                piece = _sec(sec) + big_units[unit_i]
                if pending_zero and parts:
                    parts.append("零")
                parts.append(piece)
                pending_zero = False
            elif parts:
                pending_zero = True
            unit_i += 1
        int_part = "".join(reversed(parts))
        int_part = re.sub(r"零+", "零", int_part).strip("零") or "零"

    result = int_part + "元"
    if jiao == 0 and fen == 0:
        result += "整"
    else:
        if jiao:
            result += digits[jiao] + dec_units[0]
        else:
            result += "零"
        if fen:
            result += digits[fen] + dec_units[1]
    # 第纳尔场景可在文末加「第纳尔」，模板本身已有「第纳尔」字样时由模板承载
    return result


def validate_amount_case(fields_map: Dict[str, str]) -> Optional[str]:
    lower = (
        fields_map.get("amount_lower")
        or fields_map.get("total_price_lower")
        or fields_map.get("quote_amount_digits")
        or fields_map.get("price_lower")
        or ""
    )
    upper = (
        fields_map.get("amount_upper")
        or fields_map.get("total_price_upper")
        or fields_map.get("quote_amount_words")
        or fields_map.get("price_upper")
        or ""
    )
    if not lower or not upper:
        return "请填写工程总价小写与大写"
    try:
        expect = amount_to_cn_upper(lower)
    except ValueError as e:
        return str(e)

    def norm(x: str) -> str:
        return re.sub(r"\s+", "", x).replace("整", "")

    if norm(upper) != norm(expect) and expect not in upper and upper not in expect:
        return f"大小写金额不一致：小写 {lower} 对应大写应为「{expect}」，当前为「{upper}」"
    return None


def _amount_keys_lower() -> Tuple[str, ...]:
    return ("amount_lower", "total_price_lower", "quote_amount_digits", "price_lower")


def _amount_keys_upper() -> Tuple[str, ...]:
    return ("amount_upper", "total_price_upper", "quote_amount_words", "price_upper")


def save_attachment_fields(project_id: int, code: str, fields: List[Dict[str, Any]]) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError(f"未找到附件 {code}")

    # 合并 value
    by_key = {str(f.get("key")): f for f in fields if isinstance(f, dict) and f.get("key")}
    new_fields = []
    for f in (att.get("form_data") or {}).get("fields") or []:
        item = dict(f)
        if item.get("key") in by_key:
            item["value"] = by_key[item["key"]].get("value", "")
        new_fields.append(item)
    # 若尚未分析过，直接用提交的 fields
    if not new_fields:
        new_fields = [dict(f) for f in fields]

    # 承诺函：小写金额变更时自动同步大写（可再手工改）
    by_key2 = {str(f.get("key")): f for f in new_fields}
    lower_f = next((by_key2[k] for k in _amount_keys_lower() if k in by_key2), None)
    if lower_f and lower_f.get("value"):
        try:
            upper = amount_to_cn_upper(lower_f["value"])
            updated = False
            for uk in _amount_keys_upper():
                if uk in by_key2:
                    by_key2[uk]["value"] = upper
                    updated = True
            if not updated:
                new_fields.append(
                    {
                        "key": "amount_upper",
                        "label": "工程总价（大写）",
                        "source": "quote",
                        "required": True,
                        "value": upper,
                        "hint": "由小写自动生成",
                        "evidence": "（大写）",
                    }
                )
        except Exception:
            pass

    # 「签章处」类字段改为选填（由纸质盖章完成，不必在表单填字）
    for f in new_fields:
        key = str(f.get("key") or "")
        label = str(f.get("label") or "")
        if any(x in key for x in ("seal", "signature_place", "sign_place")) or "签章处" in label or "签字处" in label:
            f["required"] = False

    form_data = {**(att.get("form_data") or {}), "fields": new_fields}
    status = att.get("status")
    if status in ("to_prepare", "pending_decision"):
        status = "drafting"
    update_instructions(project_id, {"attachment": {"code": code, "status": status, "form_data": form_data}})
    return get_attachment_workbench(project_id, code)


def complete_attachment(project_id: int, code: str) -> Dict[str, Any]:
    """校验必填后标记完成。附件可任意顺序办理，全部纳入项完成后才可进入校验。"""
    wb = get_attachment_workbench(project_id, code)
    att = wb["attachment"]
    if att.get("kind") == "power_of_attorney":
        cond = wb.get("conditions") or {}
        if cond.get("auth_filing") == "unknown":
            raise ValueError("授权备案状态为「无法确认」，不得完成或省略")
        if cond.get("auth_filing") == "filed":
            update_instructions(project_id, {"attachment": {"code": code, "status": "not_required"}})
            return get_attachment_workbench(project_id, code)

    missing = []
    fmap = {}
    for f in (att.get("form_data") or {}).get("fields") or []:
        fmap[str(f.get("key") or "")] = str(f.get("value") or "").strip()
        if f.get("required") and not str(f.get("value") or "").strip():
            missing.append(f.get("label") or f.get("key"))
    if missing:
        raise ValueError("请先填写必填字段：" + "、".join(missing))

    if att.get("kind") == "bid_commitment":
        cond = wb.get("conditions") or {}
        if not cond.get("site_visit_confirm"):
            raise ValueError("请先确认现场考察事实（已考察 / 资料评估 / 需澄清）")
        err = validate_amount_case(fmap)
        if err:
            raise ValueError(err)

    if att.get("kind") == "personnel_plan":
        form = ensure_personnel_form_data(att)
        people = form.get("people") or []
        named = [p for p in people if str(p.get("name") or "").strip()]
        if not named:
            raise ValueError("请至少填写一名管理人员姓名（可导入资料或手工新增）")
        # 模板默认四岗位：若仍保留该岗位行，则姓名必填
        for p in people:
            pos = str(p.get("position") or "").strip()
            if pos in PERSONNEL_DEFAULT_POSITIONS and not str(p.get("name") or "").strip():
                raise ValueError(f"岗位「{pos}」尚未填写人员姓名")

    if att.get("kind") == "org_chart":
        from .org_chart import validate_org_chart_complete

        validate_org_chart_complete(att)

    if att.get("kind") in (
        "labor_plan",
        "equipment_plan",
        "schedule",
        "labor_agreement",
        "confidentiality",
    ):
        from .remaining_attachments import complete_remaining

        return complete_remaining(project_id, code)

    update_instructions(project_id, {"attachment": {"code": code, "status": "done"}})
    return get_attachment_workbench(project_id, code)


def get_attachment_workbench(project_id: int, code: Optional[str] = None) -> Dict[str, Any]:
    current = get_or_init_instructions(project_id)
    state = current["result"]
    work = current_work_attachment(state)
    if not code:
        code = (work or (state.get("attachments") or [{}])[0]).get("code")
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError("附件不存在")
    if att.get("kind") == "personnel_plan":
        ensure_personnel_form_data(att)
    master = get_org_master(project_id)
    return {
        "analysis_id": current.get("id"),
        "attachment": att,
        "current_work_code": work.get("code") if work else None,
        "locked": False,
        "lock_reason": None,
        "org_master": master,
        "conditions": state.get("conditions") or {},
        "base_data": state.get("base_data") or {},
        "attachments_progress": [
            {
                "code": a["code"],
                "title": a["title"],
                "kind": a["kind"],
                "status": a["status"],
                "status_label": STATUS_LABEL.get(a.get("status") or "", a.get("status")),
                "included": a.get("included", True),
            }
            for a in state.get("attachments") or []
        ],
    }


def _fill_template_body(template_body: str, fields_list: List[Dict[str, Any]], kind: str) -> str:
    """在原模板正文上填入用户值，不另起生成条款。"""
    # .doc OLE 抽取常为 \r 换行；统一成 \n，避免 [^\n]* 跨段吞掉签章区
    text = (template_body or "").replace("\r\n", "\n").replace("\r", "\n")
    fmap = {str(f.get("key") or ""): str(f.get("value") or "").strip() for f in fields_list}

    def g(*keys: str) -> str:
        for k in keys:
            if fmap.get(k):
                return fmap[k]
        return ""

    amount_lower_keys = set(_amount_keys_lower())
    amount_upper_keys = set(_amount_keys_upper())

    # 通用：按 evidence / 常见括号空白替换（金额留给专用逻辑，避免把「小写/大写」冲掉）
    for f in fields_list:
        val = str(f.get("value") or "").strip()
        if not val:
            continue
        evidence = str(f.get("evidence") or "").strip()
        key = str(f.get("key") or "")
        label = str(f.get("label") or "")
        if key in amount_lower_keys or key in amount_upper_keys:
            continue
        if "小写" in evidence or "大写" in evidence or "小写" in label or "大写" in label:
            continue
        if key in ("currency", "sign_date", "date", "signature_date"):
            continue
        # 1) 证据串里按原空白样式回填（下划线填字，括号保括号）
        if evidence and evidence in text:
            filled_ev = evidence
            if re.search(r"[_\s　＿—－─━]{2,}", filled_ev):
                filled_ev = re.sub(r"[_\s　＿—－─━]{2,}", val, filled_ev, count=1)
            else:
                filled_ev = re.sub(r"[（(]\s*[)）]", f"（{val}）", filled_ev)
                filled_ev = re.sub(r"[（(][^）)]*[)）]", f"（{val}）", filled_ev, count=1)
            if filled_ev != evidence:
                text = text.replace(evidence, filled_ev, 1)
        # 2) 标签后空位
        for pat in (f"{label}：", f"{label}:"):
            if pat in text and val:
                text = re.sub(re.escape(pat) + r"\s*", pat + val + " ", text, count=1)

    # 承诺函专用定位（模板固定句式）
    if kind == "bid_commitment":
        signer = g("signer_name", "legal_rep_name")
        company = g("company_name", "bidder_name")
        project = g("project_name")
        employer = g("employer_name", "employer")
        lower = g(*_amount_keys_lower())
        upper = g(*_amount_keys_upper())
        currency = g("currency") or "DA"
        sign_date = g("sign_date", "date", "signature_date") or date.today().isoformat()

        if signer:
            text = re.sub(r"我\s*[（(]\s*签字人姓名\s*[)）]", f"我（{signer}）", text)
            text = re.sub(r"我\s*[（(][^）\n]{0,30}[)）]\s*以", f"我（{signer}） 以", text, count=1)
            text = re.sub(r"我[_\s　＿—－─━]{2,}以", f"我{signer}以", text, count=1)
        if company:
            text = re.sub(r"[（(]\s*公司名称\s*[)）]", f"（{company}）", text)
            text = re.sub(r"以\s*[（(][^）\n]{0,40}[)）]\s*名义", f"以（{company}） 名义", text, count=1)
            text = re.sub(r"以[_\s　＿—－─━]{2,}名义", f"以{company}名义", text, count=1)
        if project:
            text = re.sub(r"递交了\s+的投标书", f"递交了 {project} 的投标书", text)
        if employer and "中建阿尔及利亚" not in text:
            text = text.replace("向递交了", f"向{employer}递交了")
        # 金额：保留「（小写）/（大写）」字样，只填入用户值
        if lower:
            if re.search(r"工程总价：\s*（\s*小写\s*）", text):
                text = re.sub(
                    r"工程总价：\s*（\s*小写\s*）\s*[^\r\n（]*",
                    f"工程总价：（小写）{lower} {currency}",
                    text,
                    count=1,
                )
            else:
                text = re.sub(
                    r"工程总价：\s*（\s*[\d.,]+\s*）\s*[A-Za-z]*",
                    f"工程总价：（小写）{lower} {currency}",
                    text,
                    count=1,
                )
        if upper:
            if re.search(r"（\s*大写\s*）", text):
                text = re.sub(r"（\s*大写\s*）\s*[^\r\n]*", f"（大写）{upper}", text, count=1)
            else:
                text = re.sub(
                    r"（\s*[零壹贰叁肆伍陆柒捌玖拾佰仟万亿兆元角分整]+\s*）\s*第纳尔?",
                    f"（大写）{upper}",
                    text,
                    count=1,
                )
        if sign_date:
            text = re.sub(r"日\s*期：\s*[^\r\n]*", f"日期：{sign_date}", text, count=1)
            if not re.search(r"日期：\s*\d{4}", text):
                text = re.sub(r"日期：\s*", f"日期：{sign_date}", text, count=1)

    if kind == "power_of_attorney":
        # 原文：空格下划线 + (标签) → 只改括号内，保留下划线空格与 (招标人)
        legal = g("legal_rep_name")
        bidder = g("bidder_name", "company_name")
        agent_org = g("agent_org_name", "agent_unit_name", "agent_company")
        name_id = g("agent_name_id")
        agent = g("agent_name")
        agent_id = g("agent_id_no", "agent_id_card")
        if not name_id and agent and agent_id:
            name_id = f"{agent}/{agent_id}"
        elif not name_id:
            name_id = agent
        if name_id:
            text = re.sub(r"（\s*姓名/身份证号\s*）", f"（{name_id}）", text)
            text = re.sub(r"\(\s*姓名/身份证号\s*\)", f"({name_id})", text)
        if agent_org:
            text = re.sub(r"（\s*单位名称\s*）", f"（{agent_org}）", text)
            text = re.sub(r"\(\s*单位名称\s*\)", f"({agent_org})", text)
        if bidder:
            text = re.sub(r"（\s*投标人名称\s*）", f"（{bidder}）", text)
            text = re.sub(r"\(\s*投标人名称\s*\)", f"({bidder})", text)
        if legal:
            text = re.sub(r"（\s*签字人姓名\s*）", f"（{legal}）", text)
            text = re.sub(r"\(\s*签字人姓名\s*\)", f"({legal})", text)
            text = re.sub(r"（\s*姓名\s*）", f"（{legal}）", text)
            text = re.sub(r"\(\s*姓名\s*\)", f"({legal})", text)

    return text.strip()


def _docx_from_plain_text(title: str, body: str, reminders: List[str]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(16)
    for line in (body or "").splitlines():
        doc.add_paragraph(line if line.strip() else "")
    if reminders:
        doc.add_paragraph("")
        doc.add_paragraph("【签章提醒】")
        for tip in reminders:
            doc.add_paragraph(f"· {tip}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_poa_docx_bytes(project_id: int, code: str) -> Tuple[bytes, str]:
    """兼容旧接口：统一走模板回填导出。"""
    return export_attachment_docx_bytes(project_id, code)


def export_attachment_docx_bytes(project_id: int, code: str) -> Tuple[bytes, str]:
    """承诺书/授权委托书优先在原附件上填值保版式导出；失败再回退纯文本重建。"""
    wb = get_attachment_workbench(project_id, code)
    att = wb["attachment"]
    kind = att.get("kind") or ""
    form = att.get("form_data") or {}
    fields_list = form.get("fields") or []

    if kind == "personnel_plan":
        form = ensure_personnel_form_data(att)
        return _export_personnel_docx(att, form.get("people") or [])

    if kind == "org_chart":
        from .org_chart import export_org_chart_docx

        return export_org_chart_docx(project_id, code)

    if kind in (
        "labor_plan",
        "equipment_plan",
        "schedule",
        "labor_agreement",
        "confidentiality",
    ):
        from .remaining_attachments import export_remaining_docx

        return export_remaining_docx(project_id, code)

    safe = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", att.get("title") or "附件")

    # 承诺函 / 授权委托书：必须在原附件上套打，避免 OLE 抽文本乱码重建
    if kind in ("bid_commitment", "power_of_attorney"):
        content = _export_filled_original_docx(project_id, att, fields_list, kind)
        if content:
            return content, f"{safe}_已填写待签章.docx"
        rel = att.get("rel_path") or ""
        src = absolute_doc_path(project_id, rel) if rel else None
        if src is not None and src.suffix.lower() == ".doc":
            raise ValueError(
                "无法在原 .doc 附件上套打导出。请安装 Microsoft Word 或 WPS，"
                "并在环境中执行：pip install pywin32 后重启服务。"
            )
        # 仅 .docx 原件失败时才回退简版（极少）
        return _export_commitment_poa_plaintext_fallback(
            project_id, att, form, fields_list, kind, safe
        )

    raise ValueError(f"附件类型 {kind or '?'} 暂不支持 Word 导出")


def _export_filled_original_docx(
    project_id: int,
    att: Dict[str, Any],
    fields_list: List[Dict[str, Any]],
    kind: str,
) -> Optional[bytes]:
    """尝试在招标包原附件上填值，保留版式；失败返回 None。"""
    from .template_fill import fill_original_attachment

    rel = att.get("rel_path") or ""
    if not rel:
        return None
    src = absolute_doc_path(project_id, rel)
    if not src.exists():
        return None
    return fill_original_attachment(src, fields_list, kind)


def _export_commitment_poa_plaintext_fallback(
    project_id: int,
    att: Dict[str, Any],
    form: Dict[str, Any],
    fields_list: List[Dict[str, Any]],
    kind: str,
    safe: str,
) -> Tuple[bytes, str]:
    """原件套打失败时的旧逻辑：纯文本回填后新建简版 docx。"""
    template_body = form.get("template_body") or form.get("template_preview") or ""
    if not template_body:
        template_body, _ = _extract_attachment_text(project_id, att)
    elif kind == "bid_commitment":
        fresh, _ = _extract_attachment_text(project_id, att)
        if fresh and len(fresh) > 80:
            template_body = fresh

    if kind == "bid_commitment" or (template_body and len(template_body) > 80):
        filled = _fill_template_body(template_body, fields_list, kind)
        lines = []
        for line in filled.splitlines():
            cn = sum(1 for ch in line if "\u4e00" <= ch <= "\u9fff")
            ascii_alnum = sum(1 for ch in line if ch.isascii() and ch.isalnum())
            weird = sum(1 for ch in line if ord(ch) > 127 and not ("\u4e00" <= ch <= "\u9fff"))
            if weird > max(cn, 2) and cn < 4:
                continue
            if cn >= 2 or any(x in line for x in ("DA", "第纳尔", "签字", "盖章", "日期", "总价", "小写", "大写")):
                lines.append(line)
            elif line.strip() == "":
                lines.append("")
            elif ascii_alnum > 5 and weird == 0:
                lines.append(line)
        filled = "\n".join(lines).strip()
        clean_lines = []
        started = False
        for line in filled.splitlines():
            if not started:
                if "承诺" in line or "授权" in line or "委托" in line:
                    started = True
                    clean_lines.append(line)
                continue
            if re.search(r"附件\s*0?2", line) and len(clean_lines) > 3:
                break
            rare = sum(1 for ch in line if ord(ch) > 0x9FFF or ("\u3400" <= ch <= "\u4DBF"))
            common_cn = sum(1 for ch in line if "\u4e00" <= ch <= "\u9fff")
            if rare >= 3 and rare >= common_cn:
                continue
            clean_lines.append(line)
            if re.search(r"日期\s*[：:]\s*\d{4}", line):
                break
        if clean_lines:
            filled = "\n".join(clean_lines).strip()
        if kind == "bid_commitment" and "签" not in filled:
            fmap2 = {str(f.get("key") or ""): str(f.get("value") or "").strip() for f in fields_list}
            sd = fmap2.get("sign_date") or fmap2.get("date") or date.today().isoformat()
            filled = filled.rstrip() + f"\n\n公 司 签 章：\n代表人签字：\n日期：{sd}"
        if filled:
            content = _docx_from_plain_text("", filled, form.get("sign_reminder") or [])
            return content, f"{safe}_已填写待签章.docx"

    # 授权委托书：模板过短时的兜底拼装
    fields = {(f.get("key")): str(f.get("value") or "").strip() for f in fields_list}

    def g(*keys: str) -> str:
        for k in keys:
            if fields.get(k):
                return fields[k]
        return "________"

    legal = g("legal_rep_name")
    bidder = g("bidder_name")
    agent = g("agent_name")
    agent_id = g("agent_id_no", "agent_id_card")
    employer = g("employer_name", "employer")
    project = g("project_name")
    valid = g("auth_valid_until", "valid_until", "expiry_date")
    gender = g("agent_gender")
    age = g("agent_age")
    dept = g("agent_dept")
    title = g("agent_title")
    sign_date = g("sign_date", "date") or date.today().isoformat()
    agent_org = g("agent_org_name", "agent_unit_name", "bidder_name")
    y = m = d = "____"
    m2 = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", valid)
    if m2:
        y, m, d = m2.group(1), str(int(m2.group(2))), str(int(m2.group(3)))

    body = (
        f"授权委托书\n\n"
        f"本授权委托书声明：我 {legal} 系 {bidder} 的法定代表人，"
        f"现授权委托 {agent_org} 的 {agent}/{agent_id} 为我公司代理人，"
        f"参加 {employer} 的 {project} 的投标活动。"
        f"代理人在投标、开标、评标、合同谈判过程中所签署的一切文件和处理与之有关的一切事务的处理意见和回复，我均予以承认。\n"
        f"本授权书经本公司法定代表人签字并加盖法人章后生效，有效期至 {y} 年 {m} 月 {d} 日 。\n"
        f"代理人无转委托权。特此委托。\n\n"
        f"代理人：{agent}    性别：{gender}    年龄：{age}\n"
        f"单位：{agent_org}    部门：{dept}    职务：{title}\n\n"
        f"投标人：(盖章)\n法定代表人：(签字并盖章)\n日期：{sign_date}\n"
    )
    content = _docx_from_plain_text("", body, form.get("sign_reminder") or POA_SIGN_REMINDERS)
    return content, f"{safe}_已填写待签章.docx"


IMPORT_FILL_EXCLUDE_KINDS = {"org_chart", "personnel_plan"}

FIELD_FILL_SYSTEM = """你是招投标附件填表助手。
根据用户上传资料，为给定字段清单回填 value。
规则：只填原文有依据的内容；没有则 value 置空字符串；只输出 JSON。
返回：{"fields":[{"key":"...","value":"...","evidence":"..."}]}
"""


def import_fill_attachment(
    project_id: int,
    code: str,
    file_bytes: bytes,
    filename: str,
) -> Dict[str, Any]:
    """导入文件并分析回填附件字段（不含 A03 组织图、A04 人员表）。"""
    from . import instructions_analysis as inst_svc
    from . import remaining_attachments as remain_svc
    from . import config

    current = inst_svc.get_or_init_instructions(project_id)
    state = current["result"]
    att = _find_attachment(state, code=code)
    if not att:
        raise ValueError(f"未找到附件 {code}")
    kind = att.get("kind") or ""
    if kind in IMPORT_FILL_EXCLUDE_KINDS:
        raise ValueError("组织机构图与管理人员配备计划请使用各自原有入口，不支持本导入")

    root = Path(config.UPLOAD_DIR) / f"project_{project_id}" / "attachment_imports" / code
    root.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", filename or "import.bin")
    path = root / f"{uuid.uuid4().hex[:8]}_{safe_name}"
    path.write_bytes(file_bytes)

    text, method = _extract_import_text(path)
    if len(text.strip()) < 20:
        raise ValueError(f"无法读取导入文件内容（method={method}）")

    # 确保有字段清单
    if kind in ("power_of_attorney", "bid_commitment"):
        form = dict(att.get("form_data") or {})
        fields = form.get("fields") or []
        if not fields:
            wb = analyze_attachment_fields(project_id, code)
            att = wb.get("attachment") or att
            form = dict(att.get("form_data") or {})
            fields = form.get("fields") or []
        schema = [{"key": f.get("key"), "label": f.get("label")} for f in fields if f.get("key")]
        prompt = (
            f"附件: {att.get('title')} ({code})\n文件名: {filename}\n"
            f"字段清单: {schema}\n\n原文:\n{text[:16000]}"
        )
        result = call_dashscope_json(prompt, FIELD_FILL_SYSTEM, temperature=0.1)
        by_key = {
            str(x.get("key")): x
            for x in (result.get("fields") or [])
            if isinstance(x, dict) and x.get("key")
        }
        filled = 0
        new_fields = []
        for f in fields:
            item = dict(f)
            hit = by_key.get(str(f.get("key")))
            if hit and str(hit.get("value") or "").strip():
                item["value"] = str(hit.get("value")).strip()
                if hit.get("evidence"):
                    item["evidence"] = hit.get("evidence")
                filled += 1
            new_fields.append(item)
        form["fields"] = new_fields
        form["import_meta"] = {"filename": filename, "method": method, "filled": filled}
        status = att.get("status") or "to_prepare"
        if status in ("to_prepare", "pending_decision", "pending_confirm"):
            status = "drafting"
        inst_svc.update_instructions(
            project_id,
            {"attachment": {"code": code, "form_data": form, "status": status}},
        )
        return get_attachment_workbench(project_id, code)

    # 其余计划类 / 协议类：回填 form_data.fields（若有）及常见文本字段
    form = remain_svc.ensure_form_for_kind(project_id, att, state)
    fields = form.get("fields") or []
    schema = [{"key": f.get("key"), "label": f.get("label")} for f in fields if f.get("key")]
    extra_keys = []
    if kind == "labor_plan":
        extra_keys = ["peak_headcount", "credential_notes", "notes"]
    elif kind == "equipment_plan":
        extra_keys = ["voucher_note", "notes"]
    elif kind == "schedule":
        extra_keys = ["engineer_notes", "notes"]
    elif kind in ("labor_agreement", "confidentiality"):
        extra_keys = ["notes"]

    prompt = (
        f"附件: {att.get('title')} ({code})\n文件名: {filename}\n"
        f"字段清单: {schema}\n"
        f"另可回填键: {extra_keys}\n"
        f"返回 JSON: {{\"fields\":[{{\"key\":\"\",\"value\":\"\",\"evidence\":\"\"}}],\"extras\":{{}}}}\n\n"
        f"原文:\n{text[:16000]}"
    )
    result = call_dashscope_json(prompt, FIELD_FILL_SYSTEM, temperature=0.1)
    by_key = {
        str(x.get("key")): x
        for x in (result.get("fields") or [])
        if isinstance(x, dict) and x.get("key")
    }
    filled = 0
    if fields:
        new_fields = []
        for f in fields:
            item = dict(f)
            hit = by_key.get(str(f.get("key")))
            if hit and str(hit.get("value") or "").strip():
                item["value"] = str(hit.get("value")).strip()
                filled += 1
            new_fields.append(item)
        form["fields"] = new_fields
    extras = result.get("extras") if isinstance(result.get("extras"), dict) else {}
    for k in extra_keys:
        v = extras.get(k) or by_key.get(k, {}).get("value")
        if v is not None and str(v).strip():
            form[k] = str(v).strip()
            filled += 1
    form["import_meta"] = {"filename": filename, "method": method, "filled": filled}
    return remain_svc.save_plan_workbench(project_id, code, form)

